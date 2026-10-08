#!/usr/bin/env python3
"""The IMAGE SERIES gate — read from the saved FILE (the `+gen.<slot>` tags) and `series.json`, by both
delivery gates (render_deck --gate-check and codex_delivery_gate.py), so the runtimes cannot disagree.

Blocks: a series recorded with no readable/valid plan; a generated picture with no slot; a series with
NO slot placed through slot_picture; a generated person given a real-looking name, a role/quote, or
team / testimonial / customer framing on its slide with no 'fictional' label (the user's people rule,
2026-10-03 — conservative by design: better a miss than a false block). Notes: an unplaced slot; no
series-qc.json; an unresolved OFF-SERIES outlier.
"""
from __future__ import annotations

import json
import re
import shlex
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

LABELS = ("fictional", "illustrative persona", "persona (illustrative)", "composite persona",
          "虚构", "示意人物", "虚构人物", "人物为虚构")
TEAM = re.compile(r"\b(our team|meet the team|the team behind|our founders?|meet the founders?|testimonials?|"
                  r"what our (customers|clients|users) say|customer stor(y|ies))\b"
                  r"|我们的团队|团队介绍|创始团队|客户评价|用户评价|客户说|证言|用户心声", re.I)
# A person is NAMED when a name and a role sit together as a credit line — "Daniel Okafor, Head of
# Ceramics", "Ana Ruiz | Lead Gardener", "张伟｜首席设计师" — or a quote is attributed to a name. The
# final review (2026-10-03) found the first patterns blocking ordinary copy: any two Title Case words
# read as a name, a role word ANYWHERE later on the line counted, a surname-like first character (周末,
# 夏日, 金秋, 方法) began a "name", any dash after a quote was an attribution. Now the name must be the
# WHOLE segment before a separator, and the role must END the segment after it.
_SEP = r"\s*(?:,|，|\||｜|·|—{1,2}|–|\s-\s|:|：)\s*"
_FUNC = (r"(?!(?:In|At|On|For|By|To|From|With|The|A|An|Our|Your|My|Their|This|That|These|Last|Next|"
         r"Every|Each|All|Open|Join|Meet|Why|How|What|When|Where|Who)\b)")
_UP = r"[A-ZÀ-ÖØ-Þ]"                                   # a capital, accented ones too (Élodie, Óscar)
_LO = r"[^\W\d_]"                                       # any letter (Lucía, Gómez)
_EN_NAME = (r"(?:(?:Dr|Prof|Mr|Mrs|Ms|Mx)\.?\s+)?" + _FUNC + _UP + _LO + r"{1,15}(?:\s+" + _UP + r"\.)?\s+"
            + _UP + r"(?:" + _LO + r"|['’-]){1,20}")
_EN_ROLE = (r"(?i:ceo|cto|coo|cfo|vp|founder|co-founder|director|manager|lead|head|designer|engineer|"
            r"researcher|professor|gardener|teacher|coordinator|organi[sz]er|volunteer|chef|nurse|doctor|"
            r"farmer|potter|artist|owner|partner|curator)")
_EN_TAIL = r"(?:\s+(?:of|at|for)\s+[^,，|｜·—–:：]{1,40})?"
EN_CREDIT = re.compile(r"^\s*" + _EN_NAME + _SEP + r"(?:[A-Z][\w&'’-]*\s+){0,2}" + _EN_ROLE + _EN_TAIL
                       + r"\s*(?:$|[,，|｜·—–])")
ZH_SURNAMES = ("王李张刘陈杨黄赵吴周徐孙马朱胡郭何林罗高郑梁谢宋唐许韩冯邓曹彭曾肖田董袁潘蒋蔡余杜叶程苏魏吕"
               "丁任沈姚卢姜崔钟谭陆汪范金石廖贾夏韦方白邹孟熊秦邱江尹薛段雷侯龙史陶黎贺顾毛郝龚邵万钱严武戴莫孔汤")
_ZH_NAME = r"(?<![一-鿿])[" + ZH_SURNAMES + r"][一-鿿]{1,2}"
_ZH_ROLE = r"(?:首席|总监|经理|负责人|设计师|工程师|研究员|教授|创始人|主任|老师|园丁|志愿者|店主|主理人)"
ZH_CREDIT = re.compile(_ZH_NAME + r"\s*(?:——|[，,｜|·—:：])\s*[^，,｜|·—\n]{0,10}?" + _ZH_ROLE + r"\s*(?:$|[，,｜|·—。])")
# an attribution is a PERSON: 1-3 capitalised words, or a surname + 1-2 characters — then the line ends
QUOTE_ATTR = re.compile(r"[\"“].{4,}?[\"”。！？!?.]\s*[\"”]?\s*(?:—{1,2}|–|-)\s*"
                        r"(?:[A-Z][a-z]+(?:\s+[A-Z][a-z]+){0,2}|[" + ZH_SURNAMES + r"][一-鿿]{1,2})"
                        r"\s*(?:$|[,，])")


_NEG = ("non-", "non ", "not ", "非", "不是", "并非")


def _label_in(label, low):
    """`label` occurs in `low` as a label — a whole word (Latin) and not negated: "non-fictional" and "非虚构"
    (non-fiction) contain "fictional" / "虚构" and are the OPPOSITE claim (P1 review, deferred)."""
    import re as _re
    for m in _re.finditer(_re.escape(label), low):
        a, b = m.start(), m.end()
        if label[:1].isascii() and ((a and low[a - 1].isalnum()) or (b < len(low) and low[b].isalnum())):
            continue                                     # inside a longer word
        if any(low[:a].endswith(n) for n in _NEG):
            continue                                     # negated
        return True
    return False


def _named(text):
    for line in re.split(r"[\n\x0b]", text):
        if EN_CREDIT.search(line) or ZH_CREDIT.search(line) or QUOTE_ATTR.search(line):
            return True
    return bool(TEAM.search(text))


def _near(root, name, depth=3):
    """The newest `name` at most `depth` folders below `root` — the documented layout writes the
    images, their manifest and the QC report to <deck>/assets/generated, two folders down."""
    hits = [p for d in range(depth + 1) for p in root.glob("/".join(["*"] * d + [name]))]
    return max(hits, key=lambda p: p.stat().st_mtime) if hits else None


def recorded_series(gates):
    """{"imagery", "plan"} for an image-led deck — or for an imagery value that is neither "series" nor
    "selective" (a misspelt switch, e.g. "image-led", must not read as "not a series deck"); else None."""
    for sec in ("design_plan", "design"):
        d = (gates or {}).get(sec) or {}
        if isinstance(d, dict) and d.get("imagery") not in (None, "", "selective"):
            return {"imagery": d.get("imagery"), "plan": d.get("image_series") or None}
    return None


def _slide_texts(slide):
    out = []
    for sh in slide.shapes:
        if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
            out.append(sh.text_frame.text)
    return "\n".join(out)


def check(pptx, rec, deck_dir):
    findings, facts = [], {"slots": 0, "placed": 0, "generated": 0}
    if not rec:
        return findings, facts
    if rec.get("imagery") != "series":
        return [("block", "UNKNOWN IMAGERY", "imagery is {!r} — it must be \"series\" (an image-led deck) or "
                 "\"selective\" (any other deck)".format(rec.get("imagery")))], facts
    try:
        return _check(pptx, rec, deck_dir, findings, facts)
    except Exception as e:                  # a series deck the checks cannot read is not "not checked"
        return [("block", "SERIES PLAN INVALID", "the image-series checks could not run: {}: {} — fix "
                 "series.json (python3 scripts/image_series.py check <series.json>)".format(type(e).__name__, e))], facts


def _check(pptx, rec, deck_dir, findings, facts):
    import deckkit as dk
    import image_series as ims
    from pptx import Presentation
    if not rec.get("plan"):
        return [("block", "SERIES PLAN MISSING", "imagery is 'series' but design_plan.image_series names no "
                 "series.json — record the plan path")], facts
    pp = Path(rec["plan"])
    if not pp.is_absolute():
        pp = Path(deck_dir) / pp
    try:
        plan = ims.load(pp)
    except ValueError as e:
        return [("block", "SERIES PLAN MISSING", str(e))], facts
    probs = ims.check(plan)
    if probs:
        return [("block", "SERIES PLAN INVALID", "; ".join(probs[:4]) + (" …" if len(probs) > 4 else ""))], facts
    slots = {s["id"]: s for s in plan["slots"]}
    facts["slots"] = len(slots)
    placed = {}
    prs = Presentation(str(pptx))
    for n, slide in enumerate(prs.slides, 1):
        gens = []
        for sh in slide.shapes:
            sid = dk.generated_slot(sh)
            if sid is None:
                continue
            facts["generated"] += 1
            if sid not in slots:
                findings.append(("block", "UNPLANNED GENERATED IMAGE", "slide {}: a generated picture tagged "
                                 "{!r} has no slot in series.json — plan it (with its meaning line) or remove it"
                                 .format(n, sid)))
                continue
            if sid in placed and placed[sid] != n:
                findings.append(("note", "SLOT PLACED TWICE", "slot {!r} is placed on slides {} and {} — a series "
                                 "image says one thing once; give the second page its own slot".format(sid, placed[sid], n)))
            elif slots[sid].get("slide") not in (None, n):
                findings.append(("note", "SLOT ON ANOTHER SLIDE", "slot {!r} is planned for slide {} and placed on "
                                 "slide {} — fine if the deck was reordered; update the plan's slide so its file "
                                 "name and meaning follow".format(sid, slots[sid]["slide"], n)))
            placed.setdefault(sid, n)
            gens.append(slots[sid])
        people = [s for s in gens if s.get("kind") in ("generic-person", "persona")]
        if people:
            txt = _slide_texts(slide)
            low = txt.lower()
            labels = [l.lower() for l in LABELS] + [str(s.get("persona_label", "")).lower() for s in people
                                                    if s.get("persona_label")]
            labelled = any(l and _label_in(l, low) for l in labels)
            named = _named(txt)
            unlabelled_persona = any(s.get("kind") == "persona" for s in people) and not labelled
            if (named and not labelled) or unlabelled_persona:
                findings.append(("block", "GENERATED PERSON NAMED", "slide {}: a GENERATED person ({}) sits "
                                 "beside a name/role/quote or team/testimonial framing with no visible "
                                 "'fictional' label — real people get a real photo; a persona is labelled "
                                 "on its slide".format(n, ", ".join(s["id"] for s in people))))
    # a series image placed WITHOUT slot_picture has no tag, so nothing above (the people rule above all)
    # could see it — recognise the file itself: picture() embeds the image's bytes unchanged
    import hashlib
    dirs = {pp.parent}
    man = _near(pp.parent, "image_prompt_manifest.json")
    if man is not None:
        dirs.add(man.parent)
    known = {}
    for sid, sl in slots.items():
        stem = "slide-{:02d}-{}".format(sl.get("slide", 0), sid)
        for d in dirs:
            for suffix in (".png", ".cut.png", ".cut.sticker.png"):
                f = Path(d) / (stem + suffix)
                if f.is_file():
                    known[hashlib.sha1(f.read_bytes()).hexdigest()] = sid
    if known:
        for n, slide in enumerate(prs.slides, 1):
            for sh in slide.shapes:
                if dk.generated_slot(sh) is not None or not hasattr(sh, "image"):
                    continue
                try:
                    h = hashlib.sha1(sh.image.blob).hexdigest()
                except Exception:
                    continue
                if h in known:
                    findings.append(("block", "UNTAGGED SERIES IMAGE", "slide {}: slot {!r}'s image was placed "
                                     "without image_series.slot_picture, so the gate cannot apply the people rule to "
                                     "it — place it with slot_picture (sticker=True for a die-cut cut-out)"
                                     .format(n, known[h])))
    facts["placed"] = len(placed)
    if not placed:
        findings.append(("block", "NO SLOT PLACED", "imagery is 'series' but no picture was placed through "
                         "image_series.slot_picture — place each slot with it so the deck carries its tags"))
    for sid in slots:
        if sid not in placed and placed:
            findings.append(("note", "SLOT NOT PLACED", "slot {!r} (slide {}) was planned but not placed"
                             .format(sid, slots[sid]["slide"])))
    qc = _near(pp.parent, "series-qc.json")
    if qc is None:
        man = _near(pp.parent, "image_prompt_manifest.json")
        findings.append(("note", "SERIES QC MISSING", "no series-qc.json — run: python3 scripts/image_series.py "
                         "qc {} --dir {}".format(shlex.quote(str(pp)), shlex.quote(str(man.parent)) if man else
                                                 "<the folder the images were generated into>")))
    else:
        rep = json.loads(qc.read_text(encoding="utf-8"))
        ack = rep.get("acknowledged") or {}
        _qt = qc.stat().st_mtime
        newer = [Path(r.get("file", "")).name for r in rep.get("slots") or []
                 if r.get("file") and Path(r["file"]).is_file() and Path(r["file"]).stat().st_mtime > _qt]
        if newer:
            findings.append(("note", "SERIES QC STALE", "{} changed after series-qc.json was written — the report "
                             "judged older images; rerun qc".format(", ".join(newer))))
        for sid in rep.get("outliers") or []:
            if sid not in ack:
                findings.append(("note", "SERIES QC OUTLIER", "slot {!r} is OFF-SERIES — regenerate it with "
                                 "--style-ref <key>, or record why in series-qc.json 'acknowledged'".format(sid)))
    return findings, facts
