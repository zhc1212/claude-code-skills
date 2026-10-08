#!/usr/bin/env python3
"""The VISUAL LANGUAGE gate — read from the saved FILE, shared by both delivery gates
(render_deck --gate-check and codex_delivery_gate.py), so the runtimes cannot disagree.

A deck that records `design_plan.visual_language` (Codex: `design.visual_language`) must have been BUILT
in it: the cover and at least half the pages carry the language's `+vl.<name>` tag (the page functions
of visual_languages tag every shape they place), and some run uses the language's display face for the
recorded font range. "Picked collage, built plain" blocks. The language's own prohibitions (`forbids`)
are enforced here too — the register guard alone does not import library kits in a fresh process.

Blocks: UNKNOWN VISUAL LANGUAGE, LANGUAGE NOT APPLIED, DISPLAY FACE MISSING, FORBIDDEN BY LANGUAGE, and
any failure to read a deck that records a language (never a quiet NOT CHECKED).
"""
from __future__ import annotations

import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))


def recorded_language(gates):
    for sec in ("design_plan", "design"):
        d = (gates or {}).get(sec) or {}
        if isinstance(d, dict) and d.get("visual_language") not in (None, ""):
            return {"name": d.get("visual_language"), "fonts": d.get("vl_fonts") or "both",
                    "ground": d.get("vl_ground") or None}
    return None


def library_kit(gates, name):
    """The note for a register that IS the deck's recorded curated visual language, else None. The register
    notes on both runtimes ask it first: the record this skill prescribes ("bespoke <name> for …") otherwise
    reads as an INVENTED register, and both runtimes advised scaffolding a kit and saving it to the library
    for a look whose kit ships with the skill (real-deck gate-check, 2026-10-03)."""
    rec = recorded_language(gates)
    if not rec or not name or str(rec.get("name") or "").strip().lower() != str(name).strip().lower():
        return None
    try:
        import visual_languages as vl
    except Exception:
        return None
    if str(name).strip().lower() not in vl.LANGS:
        return None
    return ("`{}` is a curated visual language — its kit ships with the skill (scripts/visual_languages.py), "
            "so the kit contracts and its own prohibitions apply; nothing to scaffold or keep".format(name))


def check(pptx, rec):
    findings, facts = [], {"slides": 0, "tagged": 0}
    if not rec:
        return findings, facts
    try:
        return _check(pptx, rec, findings, facts)
    except Exception as e:                       # a recorded language the checks cannot read is not "clean"
        return [("block", "LANGUAGE NOT APPLIED", "the visual-language checks could not run on {}: {}: {}"
                 .format(Path(str(pptx)).name, type(e).__name__, e))], facts


def _check(pptx, rec, findings, facts):
    import deckkit as dk
    import visual_languages as vl
    from pptx import Presentation
    from pptx.oxml.ns import qn
    name = str(rec.get("name") or "").strip().lower()
    if name not in vl.LANGS:
        return [("block", "UNKNOWN VISUAL LANGUAGE", "design_plan.visual_language is {!r} — one of {}"
                 .format(rec.get("name"), sorted(vl.LANGS)))], facts
    ground = rec.get("ground")
    if ground not in (None, "") and ground not in vl.VARIANTS[name]:
        return [("block", "UNKNOWN GROUND VARIANT", "vl_ground is {!r} — {} has {}".format(
            ground, name, sorted(vl.VARIANTS[name])))], facts
    fonts = rec.get("fonts") if rec.get("fonts") in ("both", "mac") else "both"
    L = vl.LANGS[name]
    want = {L["fonts"]["both"]["display"]}
    if fonts == "mac":
        want.add(L["fonts"]["mac"].get("display", L["fonts"]["both"]["display"]))
    prs = Presentation(str(pptx))
    slides = list(prs.slides)
    facts["slides"] = len(slides)
    tagged, faces = [], set()
    for n, slide in enumerate(slides, 1):
        # built with a page function (tagged shapes), or started with the kit's new_slide() — an ordinary page
        # (agenda, chart) in the language's ground, which the reference sanctions; a plain add_slide is neither
        if (any(dk.vl_name(sh) == name for sh in slide.shapes)
                or slide._element.cSld.get("name") == "vl." + name):
            tagged.append(n)
        for el in slide._element.iter(qn("a:latin")):
            faces.add(el.get("typeface"))
    facts["tagged"] = len(tagged)
    cover_built = any(dk.vl_name(sh) == name for sh in slides[0].shapes) if slides else False
    if not slides or not cover_built or len(tagged) * 2 < len(slides):
        findings.append(("block", "LANGUAGE NOT APPLIED", "{} records the {!r} language but {} of {} slide(s) were "
                         "built in it{} — build the cover with the kit's cover() and at least half the pages "
                         "with visual_languages.use({!r}, prs): its page functions (cover, section, image_text, "
                         "quote, data, closing; points on the native four), or ordinary pages started with k.new_slide()"
                         .format(Path(str(pptx)).name, name, len(tagged), len(slides),
                                 "" if cover_built else " and the cover was not built with cover()", name)))
    if not (faces & want):
        findings.append(("block", "DISPLAY FACE MISSING", "no text uses {}'s display face {} for fonts={!r} — the "
                         "look did not reach the deck".format(name, sorted(want), fonts)))
    import check_register_guard as crg
    violations, _gfacts = crg.check(str(pptx), register=name)
    for v in violations:
        findings.append(("block", "FORBIDDEN BY LANGUAGE", "{} forbids it: {}".format(name, v)))
    return findings, facts
