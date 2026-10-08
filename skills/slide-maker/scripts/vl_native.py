#!/usr/bin/env python3
"""vl_native — page compositions of the NATIVE visual languages (ink, poster, cutpaper, drafting).

Imported by visual_languages only for those languages; the image-led four never load it. Each page PLANS its text
first (measured: the kit's _flow for horizontal text, vcol for vertical CJK, display for poster-scale type —
shrinking to each field's floor and refusing with VLTextOverflow past it), then draws its native art with that
text kept clear, then sets the text — so art never sits over words and z-order is art below type.
Nothing here invents a word: seal/highlight/project/icons come from the caller or are not drawn.
"""
from __future__ import annotations

import copy
from pathlib import Path

import deckkit as dk
import native_art as na
import visual_languages as vl

COMPOSERS = {}     # language -> {page: fn(k, slide, fields, image) -> rects}
GROUNDS = {}       # language -> fn(k, slide): what every page of the language carries (new_slide)


ALTS = {}          # (language, page) -> how many layouts the page can try before it refuses
_ALT = [0]         # the layout being tried now (0: the designed one); read with alt()


def register(lang, page, alts=1):
    """Register a page composition. `alts` > 1: the composition reads alt() and offers that many layouts — the
    designed one first, then roomier ones; compose() tries them in order and refuses only when the last overflows."""
    def deco(fn):
        COMPOSERS.setdefault(lang, {})[page] = fn
        ALTS[(lang, page)] = alts
        return fn
    return deco


def alt():
    return _ALT[0]


def ctx(k):
    W, H = vl._canvas(k)
    return W, H, min(W, H) / 7.5, ("land" if W >= H * 1.2 else "port")


def fit_circle(k, cx, cy, d):
    """The largest diameter <= d whose circle at (cx, cy) stays whole on the page (a margin of 0.1in at 7.5in) — a
    figure sized for one aspect ran off a square or A4 page (2026-10-05)."""
    W, H, s, _o = ctx(k)
    m = 0.1 * s
    return max(0.2, min(d, 2 * (cx - m), 2 * (W - cx - m), 2 * (cy - m), 2 * (H - cy - m)))


def text_of(fields, name):
    v = fields.get(name)
    return str(v).strip() if v is not None and str(v).strip() else None


def points_of(items):
    """2-4 points: strings, (head, line) pairs or {"head", "line"} dicts -> [(head, line or None)]."""
    if not isinstance(items, (list, tuple)) or not 2 <= len(items) <= 4:
        raise ValueError("points(): items= takes 2 to 4 points, got {!r}".format(items))
    out = []
    for it in items:
        if isinstance(it, dict):
            head, line = it.get("head"), it.get("line")
        elif isinstance(it, (list, tuple)):
            head, line = (list(it) + [None, None])[:2]
        else:
            head, line = it, None
        head = str(head or "").strip()
        if not head:
            raise ValueError("points(): every point needs its words — an empty one in {!r}".format(items))
        out.append((head, str(line).strip() if line is not None and str(line).strip() else None))
    return out


def _cjk(ch):
    o = ord(ch)
    return 0x3000 <= o <= 0x30FF or 0x3400 <= o <= 0x9FFF or 0xF900 <= o <= 0xFAFF or 0xFF00 <= o <= 0xFFEF or ch in "·—…"


def is_vertical(text):
    """Vertical only for CJK (Han/kana) with no Latin letters or digits — those would lie on their side — and only in
    a page's designed layout: when the vertical columns cannot hold the words, the page's second layout sets them
    horizontally (a 30-character title needs 4 columns; it is set across, not refused)."""
    return alt() == 0 and is_cjk_text(text)


def is_cjk_text(text):
    """Chinese or Japanese (Han/kana) with no Latin letters or digits — the text vertical setting can carry."""
    t = "".join(ch for ch in (text or "") if not ch.isspace())
    if not t or any(ch.isascii() and ch.isalnum() for ch in t):
        return False
    return dk.script_of(t) in ("han", "kana") and all(_cjk(ch) for ch in t)


def _kit_on(k, ink=None, accent=None, mute=None):
    if not (ink or accent or mute):
        return k
    kk = copy.copy(k)
    kk.P = dict(k.P)
    if ink:
        kk.P["ink"] = ink
    if mute:
        kk.P["mute"] = mute
    if accent:
        kk.P["text_accents"] = [accent] + list(k.P["text_accents"][1:])
    return kk


def flow(k, slide, page, col, items, *, anchor="top", align="l", ink=None, accent=None, mute=None, start=None):
    """PLAN horizontal text with the kit's measured flow; returns (rects, draw). ink/accent/mute override the
    colours for text that sits on a card or panel instead of the ground."""
    if not items:
        return {}, (lambda: None)
    kk = _kit_on(k, ink, accent, mute)
    return vl._flow(kk, slide, page, col, items, anchor=anchor, align=align, start=start)


def _color(k, ckey):
    return k.color("text_accents") if ckey == "accent" else k.color("mute") if ckey == "mute" else k.color("ink")


_CLAUSE_MARKS = "，。、；：！？"


def _clauses(text):
    """The text cut after each clause mark: "宋代点茶：一盏茶里的审美" -> ["宋代点茶：", "一盏茶里的审美"]."""
    out, cur = [], ""
    for ch in text:
        cur += ch
        if ch in _CLAUSE_MARKS:
            out.append(cur)
            cur = ""
    if cur:
        out.append(cur)
    return out


def vcol(k, slide, right, top, h_max, text, field, *, color=None, spacing=0.12, max_cols=2, size=None):
    """PLAN CJK `text` set vertically, right edge at `right`: starts at the field's size, shrinks toward its floor,
    refuses past `max_cols` columns at the floor. Columns break AFTER a clause mark when the clauses fit (down to 75%
    of the size) — "宋代点茶： / 一盏茶里的审美", never "宋代点茶：一 / 盏茶里的审美" (a weak-model run, 2026-10-05).
    `size=` sets the size (a couplet's two lines match) and refuses when the text needs more than max_cols columns at
    it. Returns ((x, y, w, h), size, draw)."""
    base, role, bold, ckey, _italic, floor = vl.TYPE[k.name][field]
    s = ctx(k)[2]
    fl = max(9.0, floor * s)
    n = len(text)

    def per_col(sz_):
        return max(1, int((h_max - 0.08) // (sz_ * (1.0 + spacing) / 72.0)))
    if size is not None:
        sz, cols = size, -(-n // per_col(size))
    else:
        # the FEWEST columns first: one tall column at >= 85% of the size reads as a scroll; two short ones at full
        # size read as a block (the approved cover). Only the last allowance shrinks all the way to the floor.
        for target in range(1, max_cols + 1):
            sz = base * s
            lo = fl if target == max_cols else max(fl, 0.85 * base * s)
            while True:
                cols = -(-n // per_col(sz))
                if cols <= target or sz <= lo + 1e-6:
                    break
                sz = max(lo, sz * 0.96)
            if cols <= target:
                break
    if cols > max_cols:
        raise vl.VLTextOverflow("{}: the {} {!r} needs {} vertical columns even at {:.0f}pt — shorten it".format(
            k.name, field, text[:24], cols, sz))
    paras = [text]
    clauses = _clauses(text)
    if cols >= 2 and len(clauses) >= 2:
        z = sz
        while z >= max(fl, 0.75 * sz) - 1e-6:           # pack whole clauses into columns, a little smaller if need be
            per, segs = per_col(z), [""]
            for c_ in clauses:
                if len(c_) > per:
                    segs = None
                    break
                if len(segs[-1]) + len(c_) <= per:
                    segs[-1] += c_
                else:
                    segs.append(c_)
            if segs and len(segs) <= max_cols and (size is None or z == size):
                paras, sz, cols = segs, z, len(segs)
                break
            if size is not None:
                break
            z *= 0.96
    adv = sz * (1.0 + spacing) / 72.0
    w = cols * sz * 1.28 / 72.0 + 0.06
    rows = max(len(p_) for p_ in paras) if len(paras) > 1 else -(-n // cols)     # characters in the longest column
    h = min(h_max, rows * adv + 0.12)
    x = right - w
    col = color or _color(k, ckey)

    def draw():
        tb = dk.text(slide, x, top, w, h, [k.runs(p_, sz, col, bold, role) for p_ in paras], space_after=0)
        for p in tb.text_frame.paragraphs:
            for r in p.runs:
                r._r.get_or_add_rPr().set("spc", str(int(round(sz * spacing * 100))))
        na.vertical(tb)
    return (x, top, w, h), sz, draw


def display(k, slide, rect, text, field, *, highlight=None, caps=True, ink=None, hl=None, hl_ink=None,
            align="l", anchor="t", floor=None):
    """PLAN poster-scale type in `rect`: measured, shrinking to the field's floor, refused past it; Latin set in
    capitals; `highlight` (words of the caller's own text) set on a highlighter. Lines break as the rest of the kit
    breaks them (_flow's order): at a clause mark first, else a little smaller, else a balanced narrower measure —
    never a lone CJK character or a lone word on the last line ("一盏茶的时 / 间", seen on the first render).
    Returns (rect, size, draw)."""
    base, role, bold, ckey, _italic, field_floor = vl.TYPE[k.name][field]
    floor = field_floor if floor is None else floor    # a page may let display type shrink further than its field
    s = ctx(k)[2]
    t = text.upper() if caps and not dk._has_cjk(text) else text
    hi = None
    if highlight is not None:
        hi = str(highlight).strip()
        if not hi or hi.upper() not in t.upper():
            raise ValueError("{}: highlight= must be words of the {} itself, got {!r}".format(k.name, field, highlight))
    x, y, w, h = rect
    face = k.ea_face(role, t) or k.face(role)
    sz, fl = base * s, max(9.0, floor * s)
    # measure the weight that RENDERS: Kit.run sets CJK display type bold where the Latin face has no CJK (poster's
    # Impact), and bold Hiragino is wider — measured light, "预算 budget" was one line and rendered two (2026-10-05)
    bold = bool(bold) or (dk._has_cjk(t) and role in ("display", "numeral") and bool(k.L.get("ea_heavy")))

    def measure(lines_, sz_, w_):
        # reserve what the delivery lint measures (~1.2 em a line; it never credits spacing below single), though
        # the type is SET at 0.86 — measured 2026-10-05: a block reserved at 0.86 read as overlapping the line under it
        return sum(dk.measure_text([(l_, bool(bold))], w_, sz_, font=face, line_h_factor=dk._LINT_LINE_H)
                   for l_ in (lines_ or [t]))
    import display_type as _dt

    latin_face = k.face(role)

    def words_fit(sz_):
        # every Latin word must fit the line: LibreOffice breaks "1,250,000" after a comma, and a long token inside
        # Chinese text ("预算 Supercalifragilistic…", a URL) ran 8.9in past its box with no gate noticing (final
        # review) — so each non-CJK token is measured in the Latin face it renders in, whatever else the text holds
        for wd in t.split():
            if dk._has_cjk(wd):
                continue                                   # CJK breaks between any two characters
            gw = _dt._glyph_width(wd, sz_, latin_face, bold)
            if gw is not None and gw > w - dk.TEXT_INSET_LR:
                return False
        return True
    while True:
        need = measure(None, sz, w)
        if (need <= h and words_fit(sz)) or sz <= fl + 1e-6:
            break
        sz = max(fl, sz * 0.93)
    if need > h + 1e-6 or not words_fit(sz):
        raise vl.VLTextOverflow("{}: the {} {!r} does not fit {:.2f}x{:.2f}in even at {:.0f}pt — shorten it".format(
            k.name, field, text[:24], w, h, sz))
    def widowed(sz_, w_):
        # a Latin headline stacked one word a line ("THREE / MOVES.") is a poster's stack, not a widow
        if not vl._widowed(k, field, t, sz_, w_):
            return False
        ls = vl._break_lines(k, field, t, sz_, w_)
        return dk._has_cjk(t) or len(ls) < 2 or len(ls[-2].split()) > 1

    lines, bw = None, w
    if field in vl._NO_WIDOW:
        sz2, tries = sz, 0
        while widowed(sz2, w) and sz2 * 0.95 >= fl and tries < 10:
            sz2, tries = sz2 * 0.95, tries + 1
        fit = vl._phrase_lines(k, field, t, sz, w, max(fl, min(0.7 * sz, sz2)))
        if fit is not None and (hi is None or any(hi.upper() in l_.upper() for l_ in fit[1])):
            sz2, lines = fit                       # a highlight is never split across a clause break
        elif widowed(sz2, w):
            sz2 = sz
            bw = vl._balanced_width(k, field, t, sz, w) or w
        if measure(lines, sz2, bw) <= h + 1e-6:
            sz, need = sz2, measure(lines, sz2, bw)
        else:                                      # a guard, not a path: the plain fit stands
            lines, bw = None, w
    if align == "r":
        x = x + w - bw
    elif align == "c":
        x = x + (w - bw) / 2.0
    ty = y if anchor == "t" else (y + h - need if anchor == "b" else y + (h - need) / 2.0)
    col = dk.RGBColor.from_string(ink) if ink else _color(k, ckey)

    def runs_for(piece):
        if hi and hi.upper() in piece.upper():
            i = piece.upper().index(hi.upper())
            j = i + len(hi)
            out = []
            for part, marked in ((piece[:i], False), (piece[i:j], True), (piece[j:], False)):
                if part and marked:
                    out += [dk.mark(r, hl) for r in k.runs(part, sz, dk.RGBColor.from_string(hl_ink), bold, role)]
                elif part:
                    out += k.runs(part, sz, col, bold, role)
            return out
        return k.runs(piece, sz, col, bold, role)

    def draw():
        al = {"l": dk.PP_ALIGN.LEFT, "c": dk.PP_ALIGN.CENTER, "r": dk.PP_ALIGN.RIGHT}[align]
        dk.text(slide, x, ty, bw, need, [runs_for(l_) for l_ in (lines or [t])], align=al, space_after=0,
                line_spacing=0.86)
    return (x, ty, bw, need), sz, draw


def place_image(k, slide, image, rect, page, treat="frame"):
    """The caller's picture in `rect` with the kit's own placement (feather for ink, a frame for the others)."""
    return vl._place_image(k, slide, image, rect, vl.L_((0, 0, 1, 1), treat, None), page)


def compose(k, slide, page, fields, image):
    fns = COMPOSERS.get(k.name, {})
    if page not in fns:
        raise NotImplementedError("{}.{}(): no composition is registered for this page".format(k.name, page))
    if page == "image_text" and image is None:
        raise ValueError("{}.image_text(): image= is required — for text alone use section(), quote() or points()"
                         .format(k.name))
    if not any(text_of(fields, f) for f in vl.PAGE_FIELDS[page] if f != "items") and not fields.get("items") \
            and image is None:
        raise ValueError("{}.{}(): nothing to place — pass the words".format(k.name, page))
    n0 = len(slide.shapes)
    state = dict(k.__dict__)
    last = None
    for a in range(ALTS.get((k.name, page), 1)):
        _ALT[0] = a
        try:
            rects = fns[page](k, slide, fields, image) or {}
            break
        except vl.VLTextOverflow as e:            # the copy does not fit THIS layout: undo it, try the next
            last = e
            _unbuild(slide, n0)
            k.__dict__.clear()
            k.__dict__.update(state)
        finally:
            _ALT[0] = 0
    else:
        raise last
    for sh in list(slide.shapes)[n0:]:
        dk._compose_tag(sh, vl=k.name)
    return {"rects": rects, "image": None, "free": None}


def _unbuild(slide, n0):
    """Remove everything an attempt drew after the first n0 shapes (and the picture links it added), so the next
    layout starts from the bare ground. Declarations live in shape names, so nothing else needs undoing."""
    from pptx.oxml.ns import qn
    for sh in list(slide.shapes)[n0:]:
        el = sh._element
        rids = {b.get(qn("r:embed")) for b in el.iter(qn("a:blip")) if b.get(qn("r:embed"))}
        el.getparent().remove(el)
        for rid in rids:
            try:
                slide.part.drop_rel(rid)
            except Exception:
                pass


def paint_ground(k, slide):
    fn = GROUNDS.get(k.name)
    if fn:
        n0 = len(slide.shapes)
        fn(k, slide)
        for sh in list(slide.shapes)[n0:]:
            dk._compose_tag(sh, vl=k.name)


# ═══════════════════════════════════ ink 水墨 ═══════════════════════════════════
import math as _math
import re as _re

_NUM_ZH = "一二三四"


def _ink_seal(k, slide, fields, x, y, size):
    chars = text_of(fields, "seal")
    if not chars:
        return
    W, H, _s, _o = ctx(k)
    x = min(max(x, 0.12), W - size - 0.12)
    y = min(max(y, 0.12), H - size - 0.15)
    na.seal(slide, x, y, size, chars, fill=k.P["accents"][0], ink="F6EEE6",
            face=k.ea_face("display", chars) or k.face("display"))


def _ink_sun(k, slide, cx, cy, d, clear=()):
    """The sun (a moon on the night ground) at its designed spot — or its mirror, or high on the page — wherever it
    stays clear of the words; when nowhere is clear it is left out (it is ornament; the horizontal fallback once put
    it on the closing title)."""
    W, H, s, _o = ctx(k)
    pad = 0.15 * s
    for x, y in ((cx, cy), (W - cx, cy), (cx, 0.14 * H), (W - cx, 0.14 * H)):
        r = (x - d / 2, y - d / 2, d, d)
        on_page = r[0] >= 0.1 and r[1] >= 0.1 and r[0] + d <= W - 0.1 and r[1] + d <= H - 0.1
        if on_page and not any(r[0] < c[0] + c[2] + pad and c[0] < r[0] + d + pad and r[1] < c[1] + c[3] + pad
                               and c[1] < r[1] + d + pad for c in clear):
            dk.decorative(na.disc(slide, x, y, d, k.P["sun"]), "the ink language's sun (a moon on the night ground)")
            return


def _ink_hairline(k, slide, x, y0, y1, alpha=0.45):
    return lambda: na.seg(slide, x, y0, x, y1, k.P["ink"], w=0.5, alpha=alpha)


def _run_all(draws):
    for d in draws:
        d()


@register("ink", "cover", alts=2)
def _ink_cover(k, slide, f, image):
    W, H, s, o = ctx(k)
    title, kicker, sub = text_of(f, "title"), text_of(f, "kicker"), text_of(f, "subtitle")
    rects, draws, clear = {}, [], []
    if kicker:
        r, d = flow(k, slide, "cover", (0.07 * W, 0.06 * H, 0.45 * W, 0.08 * H), [("kicker", kicker)])
        rects.update(r); draws.append(d); clear.extend(r.values())
    if title and is_vertical(title) and (sub is None or is_vertical(sub)):
        r1, _z, d1 = vcol(k, slide, W * (0.90 if o == "land" else 0.92), 0.09 * H,
                          H * (0.62 if o == "land" else 0.46), title, "title")
        rects["title"] = r1; draws.append(d1); clear.append(r1)
        if sub:
            r2, _z2, d2 = vcol(k, slide, r1[0] - 0.36 * s, 0.12 * H, r1[3] * 0.85, sub, "subtitle")
            rects["subtitle"] = r2; draws.append(d2); clear.append(r2)
            draws.append(_ink_hairline(k, slide, r1[0] - 0.18 * s, 0.12 * H, max(r2[1] + r2[3], 0.12 * H + 0.5)))
        seal_at = (r1[0] + r1[2] / 2.0, r1[1] + r1[3] + 0.2 * s)
    else:
        col = (0.50 * W, 0.16 * H, 0.42 * W, 0.40 * H) if o == "land" else (0.08 * W, 0.12 * H, 0.84 * W, 0.34 * H)
        items = [(x_, t) for x_, t in (("title", title), ("subtitle", sub)) if t]
        r, d = flow(k, slide, "cover", col, items, anchor="top")
        rects.update(r); draws.append(d); clear.extend(r.values())
        low = max((v[1] + v[3] for v in r.values()), default=col[1])
        seal_at = (col[0] + 0.3 * s, low + 0.3 * s)
    _ink_sun(k, slide, (0.20 if o == "land" else 0.22) * W, (0.30 if o == "land" else 0.58) * H, 0.95 * s, clear=clear)
    na.ink_ridges(slide, color=k.P["ink"], layers=na.INK_LAYERS[o], seed=1,
                  peak_span=(0.0, 0.62) if o == "land" else (0.0, 1.0), keep_clear=clear)
    _run_all(draws)
    z = 0.62 * s
    _ink_seal(k, slide, f, seal_at[0] - z / 2.0, seal_at[1], z)
    return rects


@register("ink", "section", alts=2)
def _ink_section(k, slide, f, image):
    W, H, s, o = ctx(k)
    num, kicker, title = text_of(f, "number"), text_of(f, "kicker"), text_of(f, "title")
    rects, draws, clear = {}, [], []
    left = (0.10 * W, 0.16 * H, 0.38 * W, 0.56 * H) if o == "land" else (0.08 * W, 0.08 * H, 0.84 * W, 0.30 * H)
    items = [(x_, t) for x_, t in (("kicker", kicker), ("number", num)) if t]
    r, d = flow(k, slide, "section", left, items, anchor="top", ink=k.P["text_accents"][0], start={"number": 150 * s})
    rects.update(r); draws.append(d); clear.extend(r.values())
    if title:
        if is_vertical(title):
            r1, _z, d1 = vcol(k, slide, W * (0.86 if o == "land" else 0.90), (0.14 if o == "land" else 0.42) * H,
                              H * (0.60 if o == "land" else 0.42), title, "title")
        else:
            col = (0.52 * W, 0.20 * H, 0.40 * W, 0.50 * H) if o == "land" else (0.08 * W, 0.42 * H, 0.84 * W, 0.34 * H)
            rr, d1 = flow(k, slide, "section", col, [("title", title)], anchor="top")
            r1 = rr.get("title", col)
        rects["title"] = r1; draws.append(d1); clear.append(r1)
    na.ink_ridges(slide, color=k.P["ink"], layers=na.INK_LAYERS["faint"], seed=2, peak_span=(0.0, 1.0), keep_clear=clear)
    _run_all(draws)
    if rects.get("title"):
        rt, z = rects["title"], 0.5 * s
        _ink_seal(k, slide, f, rt[0] + rt[2] / 2.0 - z / 2.0, rt[1] + rt[3] + 0.18 * s, z)
    return rects


@register("ink", "image_text")
def _ink_image_text(k, slide, f, image):
    W, H, s, o = ctx(k)
    if o == "land":
        img, col = (0.04 * W, 0.08 * H, 0.52 * W, 0.84 * H), (0.62 * W, 0.16 * H, 0.32 * W, 0.68 * H)
    else:
        img, col = (0.05 * W, 0.04 * H, 0.90 * W, 0.46 * H), (0.08 * W, 0.53 * H, 0.84 * W, 0.40 * H)
    items = [(x_, text_of(f, x_)) for x_ in ("kicker", "title", "body", "caption") if text_of(f, x_)]
    r, d = flow(k, slide, "image_text", col, items, anchor="middle" if o == "land" else "top")
    place_image(k, slide, image, img, "image_text", treat="feather")
    d()
    if r:
        low = max(v[1] + v[3] for v in r.values())
        z = 0.45 * s
        _ink_seal(k, slide, f, col[0], low + 0.2 * s, z)
    return r


@register("ink", "quote", alts=2)
def _ink_quote(k, slide, f, image):
    W, H, s, o = ctx(k)
    q, attr = text_of(f, "quote"), text_of(f, "attribution")
    rects, draws, clear = {}, [], []
    if q and is_vertical(q) and (attr is None or is_vertical(attr)) and o == "land":
        parts = [p for p in _re.split(r"(?<=[，、；])", q) if p]
        cols = [parts[0], "".join(parts[1:])] if len(parts) > 1 else [q]
        right, top = 0.80 * W, 0.11 * H
        # a couplet is two EQUAL columns: one size, each line one column (a weak-model run split the second line
        # into two short columns because each line was sized on its own)
        spans = [(top + i * 0.14 * H, 0.76 * H - i * 0.14 * H) for i in range(len(cols))]
        z = min(vcol(k, slide, right, y_, h_, c, "quote", max_cols=1)[1] for (y_, h_), c in zip(spans, cols))
        for i, ((y_, h_), c) in enumerate(zip(spans, cols)):
            r1, _z, d1 = vcol(k, slide, right, y_, h_, c, "quote", max_cols=1, size=z)
            rects["quote%d" % i] = r1; draws.append(d1); clear.append(r1)
            right = r1[0] - 0.40 * s
        if attr:
            r2, _z2, d2 = vcol(k, slide, right, 0.48 * H, 0.32 * H, attr, "attribution")
            rects["attribution"] = r2; draws.append(d2); clear.append(r2)
        seal_at = (right - 0.1 * s, 0.82 * H)
    else:
        col = (0.14 * W, 0.22 * H, 0.62 * W, 0.50 * H) if o == "land" else (0.08 * W, 0.16 * H, 0.84 * W, 0.46 * H)
        items = [(x_, t) for x_, t in (("quote", q), ("attribution", attr)) if t]
        r, d = flow(k, slide, "quote", col, items, anchor="middle")
        rects.update(r); draws.append(d); clear.extend(r.values())
        low = max((v[1] + v[3] for v in r.values()), default=col[1])
        seal_at = (col[0] + 0.2 * s, low + 0.25 * s)
    _ink_sun(k, slide, 0.21 * W, (0.22 if o == "land" else 0.10) * H, 0.62 * s, clear=clear)
    na.ink_ridges(slide, color=k.P["ink"], layers=na.INK_LAYERS["faint"], seed=3,
                  peak_span=(0.0, 0.45), keep_clear=clear)
    _run_all(draws)
    _ink_seal(k, slide, f, seal_at[0] - 0.24 * s, seal_at[1], 0.48 * s)
    return rects


@register("ink", "data", alts=2)
def _ink_data(k, slide, f, image):
    W, H, s, o = ctx(k)
    num, label, note = text_of(f, "number"), text_of(f, "label"), text_of(f, "note")
    if o == "land":
        cx, cy, R = 0.27 * W, 0.50 * H, 0.33 * H
    else:
        cx, cy, R = 0.50 * W, 0.28 * H, 0.32 * W
    R = fit_circle(k, cx, cy, 2 * R + 0.1 * s) / 2.0 - 0.05 * s     # the brush's wobble stays on the page too
    rects, draws, clear = {}, [], [(cx - R, cy - R, 2 * R, 2 * R)]
    if num:
        inner = (cx - R * 0.80, cy - R * 0.70, R * 1.60, R * 1.40)   # the brush ring's interior is ~1.66R wide
        r, d = flow(k, slide, "data", inner, [("number", num)], anchor="middle", align="c")
        rects.update(r); draws.append(d)
    if o == "land" and label and is_vertical(label) and (note is None or is_vertical(note)):
        x0 = cx + R + 0.9 * s
        r1, _z, d1 = vcol(k, slide, x0 + 1.2 * s, 0.18 * H, 0.60 * H, label, "label")
        rects["label"] = r1; draws.append(d1); clear.append(r1)
        if note:
            r2, _z2, d2 = vcol(k, slide, r1[0] - 0.42 * s, 0.20 * H, 0.52 * H, note, "note")
            rects["note"] = r2; draws.append(d2); clear.append(r2)
            draws.append(_ink_hairline(k, slide, r1[0] - 0.21 * s, 0.20 * H, 0.20 * H + max(r1[3], r2[3]), 0.4))
        seal_at = (r1[0] + r1[2] / 2.0, r1[1] + r1[3] + 0.2 * s)
    else:
        hcol = ((cx + R + 0.6 * s, 0.22 * H, 0.92 * W - (cx + R + 0.6 * s), 0.56 * H) if o == "land"
                else (0.08 * W, cy + R + 0.4 * s, 0.84 * W, 0.88 * H - (cy + R + 0.4 * s)))
        items = [(x_, t) for x_, t in (("label", label), ("note", note)) if t]
        r, d = flow(k, slide, "data", hcol, items, anchor="middle" if o == "land" else "top")
        rects.update(r); draws.append(d); clear.extend(r.values())
        low = max((v[1] + v[3] for v in r.values()), default=hcol[1])
        seal_at = (hcol[0] + 0.3 * s, low + 0.25 * s)
    na.enso(slide, cx, cy, R, R * 0.17, color=k.P["ink"], seed=5)
    na.ink_ridges(slide, color=k.P["ink"], layers=na.INK_LAYERS["faint"], seed=4, peak_span=(0.4, 1.0), keep_clear=clear)
    _run_all(draws)
    _ink_seal(k, slide, f, seal_at[0] - 0.22 * s, seal_at[1], 0.45 * s)
    return rects


@register("ink", "closing", alts=2)
def _ink_closing(k, slide, f, image):
    W, H, s, o = ctx(k)
    title, line = text_of(f, "title"), text_of(f, "line")
    rects, draws, clear = {}, [], []
    if title and is_vertical(title) and (line is None or is_vertical(line)):
        r1, _z, d1 = vcol(k, slide, W * (0.62 if o == "land" else 0.70), 0.10 * H, H * (0.56 if o == "land" else 0.44),
                          title, "title")
        rects["title"] = r1; draws.append(d1); clear.append(r1)
        if line:
            r2, _z2, d2 = vcol(k, slide, r1[0] - 0.4 * s, 0.14 * H, r1[3] * 0.8, line, "line")
            rects["line"] = r2; draws.append(d2); clear.append(r2)
        seal_at = (r1[0] + r1[2] / 2.0, r1[1] + r1[3] + 0.2 * s)
    else:
        col = (0.12 * W, 0.18 * H, 0.76 * W, 0.36 * H) if o == "land" else (0.08 * W, 0.12 * H, 0.84 * W, 0.34 * H)
        items = [(x_, t) for x_, t in (("title", title), ("line", line)) if t]
        r, d = flow(k, slide, "closing", col, items, anchor="top", align="c")
        rects.update(r); draws.append(d); clear.extend(r.values())
        low = max((v[1] + v[3] for v in r.values()), default=col[1])
        seal_at = (W / 2.0, low + 0.25 * s)
    _ink_sun(k, slide, (0.80 if o == "land" else 0.75) * W, (0.20 if o == "land" else 0.52) * H, 1.1 * s, clear=clear)
    na.ink_ridges(slide, color=k.P["ink"], layers=na.INK_LAYERS[o], seed=7, peak_span=(0.2, 1.0), keep_clear=clear)
    _run_all(draws)
    _ink_seal(k, slide, f, seal_at[0] - 0.25 * s, seal_at[1], 0.5 * s)
    return rects


@register("ink", "points", alts=2)
def _ink_points(k, slide, f, image):
    W, H, s, o = ctx(k)
    title, kicker = text_of(f, "title"), text_of(f, "kicker")
    pts = points_of(f.get("items"))
    vertical_ok = (o == "land" and all(is_vertical(h) and (l is None or is_vertical(l)) for h, l in pts)
                   and (title is None or is_vertical(title)))
    rects, draws, clear = {}, [], []
    if vertical_ok:                                    # columns read right to left, as a book is
        right = 0.93 * W
        if title:
            r0, _z, d0 = vcol(k, slide, right, 0.10 * H, 0.60 * H, title, "title")
            rects["title"] = r0; draws.append(d0); clear.append(r0)
            right = r0[0] - 0.55 * s
        n = len(pts)
        step = min(2.35 * s, (right - 0.10 * W) / n)
        for i, (head, line) in enumerate(pts):
            cr_ = right - i * step
            nr, nd = flow(k, slide, "points", (cr_ - step * 0.75, 0.10 * H, step * 0.6, 0.09 * H), [("mark", _NUM_ZH[i])],
                          align="c", start={"mark": 30 * s})
            rh, _z1, dh = vcol(k, slide, cr_ - step * 0.12, 0.23 * H, 0.42 * H, head, "item_head")
            draws += [nd, dh]; clear += list(nr.values()) + [rh]
            if line:
                rl, _z2, dl = vcol(k, slide, rh[0] - 0.18 * s, 0.25 * H, 0.50 * H, line, "item_line")
                draws.append(dl); clear.append(rl)
            if i < n - 1:
                draws.append(_ink_hairline(k, slide, cr_ - step * 0.96, 0.12 * H, 0.76 * H, 0.3))
        if title:
            draws.append(_ink_hairline(k, slide, rects["title"][0] - 0.28 * s, 0.12 * H, 0.76 * H, 0.45))
        seal_at = ((rects["title"][0] + rects["title"][2] / 2.0) if title else right,
                   (rects["title"][1] + rects["title"][3] + 0.2 * s) if title else 0.82 * H)
    else:
        top = 0.10 * H
        head_items = [(x_, t) for x_, t in (("kicker", kicker), ("title", title)) if t]
        r, d = flow(k, slide, "points", (0.08 * W, top, 0.84 * W, (0.20 if alt() == 0 else 0.30) * H), head_items,
                    anchor="top")
        rects.update(r); draws.append(d); clear.extend(r.values())
        y0 = max((v[1] + v[3] for v in r.values()), default=top) + 0.35 * s
        cols = 1 if (alt() == 0 or len(pts) < 3) else 2         # long copy, 3-4 points: two columns of rows
        rows = -(-len(pts) // cols)
        row_h, colw = (0.86 * H - y0) / rows, 0.84 * W / cols
        cjk_num = title is not None and is_cjk_text(title)
        for i, (head, line) in enumerate(pts):
            x, y = 0.08 * W + (i % cols) * colw, y0 + (i // cols) * row_h
            nr, nd = flow(k, slide, "points", (x, y, 0.9 * s, row_h), [("mark", _NUM_ZH[i] if cjk_num else str(i + 1))],
                          start={"mark": 30 * s})
            tr, td_ = flow(k, slide, "points", (x + 1.0 * s, y, colw - 1.15 * s, row_h - 0.12 * s),
                           [(x_, t) for x_, t in (("item_head", head), ("item_line", line)) if t], anchor="top")
            draws += [nd, td_]; clear += list(nr.values()) + list(tr.values())
            if i // cols < rows - 1:
                yy = y + row_h - 0.06 * s
                draws.append(lambda x=x, yy=yy: na.seg(slide, x, yy, x + colw - 0.25 * s, yy, k.P["ink"], w=0.5, alpha=0.25))
        seal_at = (0.86 * W, 0.12 * H)
    na.ink_ridges(slide, color=k.P["ink"], layers=na.INK_LAYERS["faint"], seed=6, peak_span=(0.0, 1.0), keep_clear=clear)
    _run_all(draws)
    _ink_seal(k, slide, f, seal_at[0] - 0.25 * s, seal_at[1], 0.5 * s)
    return rects


# ═══════════════════════════════════ poster 海报大字 ═══════════════════════════════════
# One saturated field per page, in turn. Every text colour is chosen against ITS field (tested >= 4.5:1).
POSTER_FIELDS = {
    "light": [
        dict(bg="1F3BFF", ink="F3F0E8", accent="D7FF3B", panel="141414", panel_ink="F3F0E8", panel_accent="D7FF3B",
             hl="D7FF3B", hl_ink="141414", block=("D7FF3B", "141414")),
        dict(bg="D7FF3B", ink="141414", accent="1F3BFF", panel="141414", panel_ink="F3F0E8", panel_accent="D7FF3B",
             hl="141414", hl_ink="D7FF3B", block=("1F3BFF", "141414")),
        dict(bg="141414", ink="F3F0E8", accent="D7FF3B", panel="D7FF3B", panel_ink="141414", panel_accent="141414",
             hl="D7FF3B", hl_ink="141414", block=("D7FF3B", "FF4B1F")),
        dict(bg="FF4B1F", ink="141414", accent="141414", panel="141414", panel_ink="F3F0E8", panel_accent="FF4B1F",
             hl="141414", hl_ink="FF4B1F", block=("141414", "F3F0E8"))],
    "paper": [
        dict(bg="F3F0E8", ink="141414", accent="1F3BFF", panel="141414", panel_ink="F3F0E8", panel_accent="F3F0E8",
             hl="1F3BFF", hl_ink="F3F0E8", block=("1F3BFF", "141414")),
        dict(bg="F3F0E8", ink="141414", accent="1F3BFF", panel="1F3BFF", panel_ink="F3F0E8", panel_accent="F3F0E8",
             hl="141414", hl_ink="F3F0E8", block=("141414", "1F3BFF")),
        dict(bg="141414", ink="F3F0E8", accent="9AA8FF", panel="F3F0E8", panel_ink="141414", panel_accent="1F3BFF",
             hl="1F3BFF", hl_ink="F3F0E8", block=("1F3BFF", "F3F0E8")),
        dict(bg="1F3BFF", ink="F3F0E8", accent="F3F0E8", panel="F3F0E8", panel_ink="141414", panel_accent="1F3BFF",
             hl="F3F0E8", hl_ink="1F3BFF", block=("F3F0E8", "141414"))],
}


def _away(bg, ink, t=0.14):
    """An ordinary page's card on a field: the field pushed AWAY from its ink (a near-black field is lifted), so the
    field's own ink and accent read on the card as they do on the field."""
    b = vl._rgb(bg)
    if vl._lum(ink) > vl._lum(bg):
        to, t = ((255, 255, 255), 0.10) if vl._lum(bg) < 0.03 else ((0, 0, 0), t)
    else:
        to = (255, 255, 255)
    return "".join("{:02X}".format(int(round(b[j] * (1 - t) + to[j] * t))) for j in range(3))


def _poster_ground(k, slide):
    """Every poster page takes the next colour field. The deck's default inks (dk.DEEP …), its card (rs.card) and the
    kit's palette follow the field, so an ORDINARY page started with k.new_slide() stays readable on cobalt as on
    lime — measured 2026-10-05: with the deck-wide palette, a cream dk.DEEP sat on the lime field at 1.1:1."""
    fields = POSTER_FIELDS[k.ground]
    fld = fields[(len(k.prs.slides) - 1) % len(fields)]
    k.field = fld
    card = _away(fld["bg"], fld["ink"])
    k.P = dict(k.P, ground=fld["bg"], ink=fld["ink"], mute=fld["ink"], panel=card,
               text_accents=[fld["accent"], fld["ink"]])
    vl._PAL_OVERRIDE["poster"] = k.P
    na.solid_background(slide, fld["bg"])
    dk.set_ground(fld["bg"])
    dk.set_palette(deep=fld["ink"], slate=fld["ink"], mute=fld["ink"], tint=card, magenta=fld["accent"],
                   accents=[fld["accent"], fld["ink"]])


GROUNDS["poster"] = _poster_ground


def _meta_height(k, kicker):
    """How tall the credit line's kicker is — measured, so a long one wraps and the page starts below it (a 13-word
    kicker ran into the cover title on a square page, 2026-10-05)."""
    W, H, s, _o = ctx(k)
    if not kicker:
        return 0.4 * s
    t = kicker if dk._has_cjk(kicker) else kicker.upper()
    face = k.ea_face("mono", t) or k.face("mono")
    return max(0.4 * s, dk.measure_text([(t, True)], 0.62 * W, max(9.0, 11.0 * s), font=face,
                                        line_h_factor=dk._LINT_LINE_H))


def _below_meta(k, rect, kicker):
    """`rect` starting below the credit line (shortened by what it gave up), when the kicker reaches into it."""
    W, H, s, _o = ctx(k)
    x, y, w, h = rect
    floor_y = 0.045 * H + _meta_height(k, kicker) + 0.12 * s
    if y >= floor_y:
        return rect
    return (x, floor_y, w, max(0.3, h - (floor_y - y)))


def _poster_meta(k, slide, kicker):
    """The fine print: the caller's kicker left, the page number right — a poster's credit line."""
    W, H, s, _o = ctx(k)
    fld = k.field
    sz = max(9.0, 11.0 * s)
    col = dk.RGBColor.from_string(fld["ink"])
    if kicker:
        t = kicker if dk._has_cjk(kicker) else kicker.upper()
        dk.text(slide, 0.04 * W, 0.045 * H, 0.62 * W, _meta_height(k, kicker), [k.runs(t, sz, col, True, "mono")],
                space_after=0)
    dk.text(slide, 0.76 * W, 0.045 * H, 0.20 * W, 0.4 * s, [k.runs("{:02d}".format(len(k.prs.slides)), sz, col, True, "mono")],
            align=dk.PP_ALIGN.RIGHT, space_after=0)


def _blocks(k, slide, blocks):
    for (bx, by, bw, bh, deg), fill in zip(blocks, k.field["block"]):
        na.clipped_block(slide, bx, by, bw, bh, deg, fill=fill)


@register("poster", "cover")
def _poster_cover(k, slide, f, image):
    W, H, s, o = ctx(k)
    fld = k.field
    title, kicker, sub, hl = text_of(f, "title"), text_of(f, "kicker"), text_of(f, "subtitle"), text_of(f, "highlight")
    if o == "land":
        trect = (0.035 * W, 0.11 * H, 0.56 * W, 0.84 * H)
        blocks = ((0.63 * W, 0.45 * H, 0.37 * W, 0.37 * W, -9.0), (0.79 * W, 0.66 * H, 0.26 * W, 0.39 * H, -9.0))
    else:
        trect = (0.06 * W, 0.08 * H, 0.88 * W, 0.56 * H)
        blocks = ((0.46 * W, 0.74 * H, 0.60 * W, 0.40 * W, -9.0), (0.70 * W, 0.86 * H, 0.40 * W, 0.20 * H, -9.0))
    trect = _below_meta(k, trect, kicker)
    rects, draws = {}, []
    if sub:
        srect = (trect[0], trect[1] + trect[3] - 0.9 * s, trect[2], 0.9 * s)
        r, d = flow(k, slide, "cover", srect, [("subtitle", sub)], anchor="bottom")
        rects.update(r); draws.append(d)
        trect = (trect[0], trect[1], trect[2], trect[3] - 1.0 * s)
    if title:
        r1, _z, d1 = display(k, slide, trect, title, "title", highlight=hl, hl=fld["hl"], hl_ink=fld["hl_ink"],
                             floor=vl.TYPE[k.name]["label"][5])
        rects["title"] = r1; draws.append(d1)
    _blocks(k, slide, blocks)
    _poster_meta(k, slide, kicker)
    _run_all(draws)
    return rects


@register("poster", "section")
def _poster_section(k, slide, f, image):
    W, H, s, o = ctx(k)
    fld = k.field
    num, kicker, title, hl = text_of(f, "number"), text_of(f, "kicker"), text_of(f, "title"), text_of(f, "highlight")
    nrect = _below_meta(k, (0.04 * W, 0.12 * H, 0.44 * W, 0.80 * H) if o == "land" else (0.06 * W, 0.08 * H, 0.88 * W, 0.40 * H),
                        kicker)
    trect = (0.52 * W, 0.26 * H, 0.44 * W, 0.56 * H) if o == "land" else (0.06 * W, 0.52 * H, 0.88 * W, 0.36 * H)
    rects, draws = {}, []
    if num:
        r, _z, d = display(k, slide, nrect, num, "number", caps=False, ink=fld["accent"], anchor="b")
        rects["number"] = r; draws.append(d)
    if title:
        r, _z, d = display(k, slide, trect, title, "label", highlight=hl, hl=fld["hl"], hl_ink=fld["hl_ink"])
        rects["title"] = r; draws.append(d)
    _poster_meta(k, slide, kicker)
    _run_all(draws)
    return rects


@register("poster", "image_text")
def _poster_image_text(k, slide, f, image):
    W, H, s, o = ctx(k)
    img = (0.52 * W, 0.12 * H, 0.44 * W, 0.76 * H) if o == "land" else (0.06 * W, 0.08 * H, 0.88 * W, 0.42 * H)
    col = (0.05 * W, 0.14 * H, 0.42 * W, 0.72 * H) if o == "land" else (0.06 * W, 0.54 * H, 0.88 * W, 0.38 * H)
    items = [(x_, text_of(f, x_)) for x_ in ("title", "body", "caption") if text_of(f, x_)]
    r, d = flow(k, slide, "image_text", col, items, anchor="middle" if o == "land" else "top")
    place_image(k, slide, image, img, "image_text")
    _poster_meta(k, slide, text_of(f, "kicker"))
    d()
    return r


@register("poster", "quote")
def _poster_quote(k, slide, f, image):
    W, H, s, o = ctx(k)
    fld = k.field
    q, attr, hl = text_of(f, "quote"), text_of(f, "attribution"), text_of(f, "highlight")
    rects, draws = {}, []
    mrect = (0.03 * W, 0.08 * H, 0.18 * W, 0.40 * H) if o == "land" else (0.06 * W, 0.06 * H, 0.30 * W, 0.18 * H)
    qrect = (0.20 * W, 0.20 * H, 0.74 * W, 0.56 * H) if o == "land" else (0.06 * W, 0.26 * H, 0.88 * W, 0.52 * H)
    r, _z, d = display(k, slide, mrect, "“", "mark", caps=False, ink=fld["accent"], floor=36)
    draws.append(d)
    if q:
        r1, _z, d1 = display(k, slide, qrect, q, "quote", highlight=hl, hl=fld["hl"], hl_ink=fld["hl_ink"])
        rects["quote"] = r1; draws.append(d1)
        if attr:
            ar, ad = flow(k, slide, "quote", (qrect[0], r1[1] + r1[3] + 0.25 * s, qrect[2], 0.6 * s),
                          [("attribution", attr)], accent=fld["accent"])
            rects.update(ar); draws.append(ad)
    _poster_meta(k, slide, None)
    _run_all(draws)
    return rects


@register("poster", "data")
def _poster_data(k, slide, f, image):
    W, H, s, o = ctx(k)
    num, label, note = text_of(f, "number"), text_of(f, "label"), text_of(f, "note")
    nrect = (0.02 * W, 0.08 * H, 0.52 * W, 0.86 * H) if o == "land" else (0.05 * W, 0.08 * H, 0.90 * W, 0.42 * H)
    lrect = (0.56 * W, 0.24 * H, 0.40 * W, 0.34 * H) if o == "land" else (0.06 * W, 0.54 * H, 0.88 * W, 0.20 * H)
    rects, draws = {}, []
    if num:
        r, _z, d = display(k, slide, nrect, num, "number", caps=False, anchor="b", floor=48)   # 1,250,000 fits one line
        rects["number"] = r; draws.append(d)
    if label:
        r1, _z, d1 = display(k, slide, lrect, label, "label")
        rects["label"] = r1; draws.append(d1)
        if note:
            nr, nd = flow(k, slide, "data", (lrect[0], r1[1] + r1[3] + 0.2 * s, lrect[2], 1.2 * s), [("note", note)])
            rects.update(nr); draws.append(nd)
    _poster_meta(k, slide, text_of(f, "kicker"))
    _run_all(draws)
    return rects


@register("poster", "closing")
def _poster_closing(k, slide, f, image):
    W, H, s, o = ctx(k)
    fld = k.field
    title, line, hl = text_of(f, "title"), text_of(f, "line"), text_of(f, "highlight")
    trect = (0.05 * W, 0.14 * H, 0.70 * W, 0.58 * H) if o == "land" else (0.06 * W, 0.12 * H, 0.88 * W, 0.50 * H)
    rects, draws = {}, []
    if title:
        r, _z, d = display(k, slide, trect, title, "title", highlight=hl, hl=fld["hl"], hl_ink=fld["hl_ink"],
                           floor=vl.TYPE[k.name]["label"][5])
        rects["title"] = r; draws.append(d)
        if line:
            lr, ld = flow(k, slide, "closing", (trect[0], r[1] + r[3] + 0.3 * s, trect[2], 1.0 * s), [("line", line)])
            rects.update(lr); draws.append(ld)
    blocks = ((0.80 * W, 0.62 * H, 0.30 * W, 0.30 * W, 12.0),) if o == "land" else ((0.62 * W, 0.80 * H, 0.50 * W, 0.24 * H, 12.0),)
    _blocks(k, slide, blocks)
    _poster_meta(k, slide, None)
    _run_all(draws)
    return rects


@register("poster", "points", alts=2)
def _poster_points(k, slide, f, image):
    """Points as a rising staircase of panels beside the title; when the steps are too narrow for their words (a
    small canvas, long lines), the second layout puts the title across the top and the steps underneath as full-width
    bars that step in, one per point."""
    W, H, s, o = ctx(k)
    fld = k.field
    title, kicker = text_of(f, "title"), text_of(f, "kicker")
    pts = points_of(f.get("items"))
    n = len(pts)
    rects, draws = {}, []
    gap = 0.15 * s
    if alt() == 0:
        if o == "land":
            trect, (rx, ry, rw, rh) = (0.04 * W, 0.11 * H, 0.42 * W, 0.42 * H), (0.50 * W, 0.30 * H, 0.46 * W, 0.63 * H)
        else:
            trect, (rx, ry, rw, rh) = (0.06 * W, 0.08 * H, 0.88 * W, 0.20 * H), (0.06 * W, 0.34 * H, 0.88 * W, 0.58 * H)
        trect = _below_meta(k, trect, kicker)
        if title:
            # display size first, but a long title shrinks to the label's floor rather than being refused (measured
            # 2026-10-05: at the title field's 44pt floor, CI's wider Linux faces refused the matrix copy on 4:3)
            r, _z, d = display(k, slide, trect, title, "title", floor=vl.TYPE[k.name]["label"][5])
            rects["title"] = r; draws.append(d)
        pw = (rw - (n - 1) * gap) / n
        for i, (head, line) in enumerate(pts):
            ph = rh * (0.64 + 0.36 * (i / float(max(n - 1, 1))))
            px, py = rx + i * (pw + gap), ry + rh - ph
            dk.box(slide, px, py, pw, ph, fill=fld["panel"])
            mh = min(1.4 * s, ph * 0.4)
            _nr, _z, nd = display(k, slide, (px + 0.12 * s, py + 0.10 * s, pw - 0.24 * s, mh), "{:02d}".format(i + 1),
                                  "mark", caps=False, ink=fld["panel_accent"], floor=12)   # the numeral fits its panel
            tr, td_ = flow(k, slide, "points", (px + 0.14 * s, py + ph * 0.45, pw - 0.28 * s, ph * 0.52),
                           [(x_, t) for x_, t in (("item_head", head), ("item_line", line)) if t], anchor="bottom",
                           ink=fld["panel_ink"], mute=fld["panel_ink"])
            draws += [nd, td_]
    else:
        # a shorter band than the designed layout's: display type fills the rect it gets, and the steps need the room
        trect = (0.04 * W, 0.09 * H, 0.92 * W, 0.17 * H) if o == "land" else (0.06 * W, 0.07 * H, 0.88 * W, 0.21 * H)
        trect = _below_meta(k, trect, kicker)
        top = trect[1]
        if title:                    # the roomier layout lets the title go smaller still (20pt) before it refuses
            r, _z, d = display(k, slide, trect, title, "title", floor=min(20, vl.TYPE[k.name]["label"][5]))
            rects["title"] = r; draws.append(d)
            top = r[1] + r[3]
        rx, ry, rw = trect[0], top + 0.25 * s, trect[2]
        rh = H * 0.93 - ry
        ph = (rh - (n - 1) * gap) / n
        ind = min(0.55 * s, rw * 0.06)
        for i, (head, line) in enumerate(pts):
            px, py, pw = rx + i * ind, ry + i * (ph + gap), rw - i * ind
            dk.box(slide, px, py, pw, ph, fill=fld["panel"])
            nw = min(1.3 * s, pw * 0.18)
            _nr, _z, nd = display(k, slide, (px + 0.15 * s, py + 0.04 * s, nw, ph - 0.08 * s), "{:02d}".format(i + 1),
                                  "mark", caps=False, ink=fld["panel_accent"], anchor="m", floor=12)
            tx = px + nw + 0.35 * s
            tr, td_ = flow(k, slide, "points", (tx, py + 0.06 * s, px + pw - 0.2 * s - tx, ph - 0.12 * s),
                           [(x_, t) for x_, t in (("item_head", head), ("item_line", line)) if t], anchor="middle",
                           ink=fld["panel_ink"], mute=fld["panel_ink"])
            draws += [nd, td_]
    _poster_meta(k, slide, kicker)
    _run_all(draws)
    return rects


# ═══════════════════════════════════ cutpaper 剪纸层叠 ═══════════════════════════════════
CUT_ART = {
    "light": {"hills": ("D3E8C9", "A2CFA3", "5AA38A", "2F6F62"), "sun": ("F7D38A", "F2B33D"),
              "rings": ("FBE6B4", "F7D38A", "F2B33D", "E8902F"), "cloud": "FFFFFF", "sheets": ("F2B33D", "EE8A6B"),
              "discs": ("C2553A", "2E6DA4", "5B3F6E", "2F7F69")},
    "night": {"hills": ("2E4A63", "3B6476", "467E7A", "234F4E"), "sun": ("D9CDA6", "F1E7C8"),
              "rings": ("39466A", "D9CDA6", "F1E7C8", "C9B98A"), "cloud": "C9D3E3", "sheets": ("F1E7C8", "EE8A6B"),
              "discs": ("C2553A", "2E6DA4", "5B3F6E", "2F7F69")},
}


def _card_flow(k, slide, page, col, items, **kw):
    return flow(k, slide, page, col, items, ink=k.P["card_ink"], mute=k.P["card_mute"], accent=k.P["card_accent"], **kw)


def _back_hills(k, slide, o, seed=1):
    W, H, s, _o = ctx(k)
    A = CUT_ART[k.ground]
    base = (0.61, 0.71) if o == "land" else (0.70, 0.78)
    na.paper_hill(slide, na.hill_points(W, base[0] * H, 0.16 * H if o == "land" else 0.08 * H, seed, 1.3), H, fill=A["hills"][0])
    na.paper_hill(slide, na.hill_points(W, base[1] * H, 0.15 * H if o == "land" else 0.07 * H, seed + 3, 1.7), H, fill=A["hills"][1])


def _front_hills(k, slide, o, seed=7, low=False):
    W, H, s, _o = ctx(k)
    A = CUT_ART[k.ground]
    if not low:
        na.paper_hill(slide, na.hill_points(W, (0.82 if o == "land" else 0.86) * H, 0.08 * H, seed, 1.9), H, fill=A["hills"][2])
    na.paper_hill(slide, na.hill_points(W, (0.92 if o == "land" else 0.93) * H, 0.05 * H, seed + 2, 2.6), H, fill=A["hills"][3])


@register("cutpaper", "cover", alts=2)
def _cut_cover(k, slide, f, image):
    W, H, s, o = ctx(k)
    A = CUT_ART[k.ground]
    if o == "land":
        card, sun, clouds = (0.064 * W, 0.14 * H, 0.48 * W, 0.47 * H), (0.76 * W, 0.23 * H, 0.35 * H), \
            ((0.56 * W, 0.17 * H, 0.15 * W), (0.91 * W, 0.43 * H, 0.11 * W))
        if alt():                    # long copy: a wider, taller card, still tucked into the hills
            card = (0.05 * W, 0.07 * H, 0.60 * W, 0.66 * H)
    else:
        card, sun, clouds = (0.08 * W, 0.10 * H, 0.84 * W, 0.32 * H), (0.70 * W, 0.53 * H, 0.26 * W), \
            ((0.28 * W, 0.50 * H, 0.30 * W),)
        if alt():                    # a column with margin for wider faces (Linux substitutes measure wider)
            card = (0.06 * W, 0.06 * H, 0.88 * W, 0.56 * H)
    pad = 0.4 * s
    items = [(x_, text_of(f, x_)) for x_ in ("kicker", "title", "subtitle") if text_of(f, x_)]
    r, d = _card_flow(k, slide, "cover", (card[0] + pad, card[1] + pad, card[2] - 2 * pad, card[3] - 2 * pad), items,
                      start={"title": 76 * s})
    text_bottom = max((v[1] + v[3] for v in r.values()), default=card[1])
    if o != "land":                  # a free-standing card fits its words (landscape keeps its height for the tuck)
        card = (card[0], card[1], card[2], text_bottom + pad - card[1])
    if o != "land":                  # under the card (fitted to its words), never buried beneath it
        foot = card[1] + card[3]
        sun = (sun[0], max(sun[1], foot + sun[2] * 0.55 + 0.1 * s), sun[2])
        clouds = tuple((cx_, max(cy_, foot + 0.22 * cw_ + 0.1 * s), cw_) for cx_, cy_, cw_ in clouds)   # by its height
    na.paper_sun(slide, sun[0], sun[1], (sun[2], sun[2] * 0.7), A["sun"])
    for cx, cy, cw in clouds:
        na.cloud(slide, cx, cy, cw, fill=A["cloud"])
    _back_hills(k, slide, o)
    na.paper_card(slide, *card)
    if o == "land":                  # the near hill tucks the card's foot into the scene — never above its words
        foot = card[1] + card[3]
        tuck = max(text_bottom + 0.15 * s, foot - 0.30 * s)
        right = card[0] + card[2]
        pts = [(W * i / 16.0, min(H - 0.3, tuck + 0.12 * s * _math.sin(i * 0.9) + max(0.0, W * i / 16.0 - right) * 0.16))
               for i in range(17)]
        near = na.paper_hill(slide, pts, H, fill=A["hills"][2])
        dk.overlap_intent(near, "the near paper hill tucks the title card's foot into the scene")
        _front_hills(k, slide, o, low=True)
    else:
        _front_hills(k, slide, o)
    d()
    return r


@register("cutpaper", "section", alts=2)
def _cut_section(k, slide, f, image):
    W, H, s, o = ctx(k)
    A = CUT_ART[k.ground]
    num, kicker, title = text_of(f, "number"), text_of(f, "kicker"), text_of(f, "title")
    cx, cy, d0 = (0.24 * W, 0.46 * H, 0.52 * H) if o == "land" else (0.50 * W, 0.26 * H, 0.56 * W)
    if alt() and o != "land":                # long copy: a smaller sun higher up, the words get the room below
        cy, d0 = 0.20 * H, 0.36 * W
    d0 = fit_circle(k, cx, cy, d0)
    rects, draws = {}, []
    if num:
        r, d = _card_flow(k, slide, "section", (cx - d0 * 0.3, cy - d0 * 0.3, d0 * 0.6, d0 * 0.6), [("number", num)],
                          anchor="middle", align="c", start={"number": 120 * s})
        rects.update(r); draws.append(d)
    col = (0.48 * W, 0.24 * H, 0.44 * W, 0.46 * H) if o == "land" else (0.08 * W, 0.56 * H, 0.84 * W, 0.28 * H)
    if alt():
        col = (0.46 * W, 0.12 * H, 0.48 * W, 0.70 * H) if o == "land" else (0.08 * W, cy + d0 / 2 + 0.3 * s, 0.84 * W,
                                                                             0.86 * H - (cy + d0 / 2 + 0.3 * s))
    items = [(x_, t) for x_, t in (("kicker", kicker), ("title", title)) if t]
    r, d = flow(k, slide, "section", col, items, anchor="middle" if o == "land" else "top")
    rects.update(r); draws.append(d)
    na.paper_sun(slide, cx, cy, (d0, d0 * 0.72), (A["rings"][1], A["rings"][3]))
    _front_hills(k, slide, o, seed=11)
    _run_all(draws)
    return rects


@register("cutpaper", "image_text")
def _cut_image_text(k, slide, f, image):
    W, H, s, o = ctx(k)
    frame = (0.06 * W, 0.10 * H, 0.46 * W, 0.76 * H) if o == "land" else (0.08 * W, 0.05 * H, 0.84 * W, 0.44 * H)
    card = (0.57 * W, 0.18 * H, 0.37 * W, 0.62 * H) if o == "land" else (0.08 * W, 0.53 * H, 0.84 * W, 0.36 * H)
    pad = 0.35 * s
    items = [(x_, text_of(f, x_)) for x_ in ("kicker", "title", "body", "caption") if text_of(f, x_)]
    r, d = _card_flow(k, slide, "image_text", (card[0] + pad, card[1] + pad, card[2] - 2 * pad, card[3] - 2 * pad), items,
                      anchor="middle")
    na.paper_card(slide, *frame)
    inset = 0.14 * s
    place_image(k, slide, image, (frame[0] + inset, frame[1] + inset, frame[2] - 2 * inset, frame[3] - 2 * inset), "image_text")
    na.paper_card(slide, *card)
    d()
    return r


@register("cutpaper", "quote")
def _cut_quote(k, slide, f, image):
    W, H, s, o = ctx(k)
    A = CUT_ART[k.ground]
    card = (0.12 * W, 0.17 * H, 0.76 * W, 0.61 * H) if o == "land" else (0.08 * W, 0.18 * H, 0.84 * W, 0.56 * H)
    # the stack SHOWS: each sheet is the card's own size, offset down and right so it peeks past the card's edges
    # (the plan's smaller sheets sat wholly behind it — measured 2026-10-05, nothing but a rotated corner showed)
    for (dx, dy, rot), col in zip(((0.42, 0.40, 2.5), (0.22, 0.20, -1.5)), A["sheets"]):
        na.paper_card(slide, card[0] + dx * s, card[1] + dy * s, card[2], card[3], fill=col, rotation=rot)
    na.paper_card(slide, *card)
    pad = 0.6 * s
    items = [(x_, text_of(f, x_)) for x_ in ("quote", "attribution") if text_of(f, x_)]
    items = [("mark", "“")] + items
    r, d = _card_flow(k, slide, "quote", (card[0] + pad, card[1] + pad * 0.6, card[2] - 2 * pad, card[3] - 1.2 * pad), items,
                      anchor="middle")
    d()
    return r


@register("cutpaper", "data")
def _cut_data(k, slide, f, image):
    W, H, s, o = ctx(k)
    A = CUT_ART[k.ground]
    num, label, note = text_of(f, "number"), text_of(f, "label"), text_of(f, "note")
    cx, cy, D = (0.27 * W, 0.50 * H, min(0.86 * H, 0.52 * W)) if o == "land" else (0.50 * W, 0.28 * H, 0.86 * W)
    cx = max(cx, D / 2.0 + 0.1 * s)              # the whole sun on the page, on any canvas (4:3 ran 0.52in off it)
    D = fit_circle(k, cx, cy, D)
    rects, draws = {}, []
    if num:
        inner, wide = D * 0.39, D * 0.72         # the figure is as tall as the inner disc; a long one spans the rings
        r, d = _card_flow(k, slide, "data", (cx - wide / 2, cy - inner / 2, wide, inner), [("number", num)],
                          anchor="middle", align="c")
        rects.update(r); draws.append(d)
    col = (cx + D / 2 + 0.5 * s, 0.24 * H, 0.94 * W - (cx + D / 2 + 0.5 * s), 0.50 * H) if o == "land" else \
        (0.08 * W, cy + D / 2 + 0.4 * s, 0.84 * W, 0.84 * H - (cy + D / 2 + 0.4 * s))
    items = [(x_, t) for x_, t in (("label", label), ("note", note)) if t]
    r, d = flow(k, slide, "data", col, items, anchor="middle" if o == "land" else "top")
    rects.update(r); draws.append(d)
    na.paper_sun(slide, cx, cy, tuple(D * x for x in (1.0, 0.78, 0.58, 0.39)), A["rings"])
    _run_all(draws)
    return rects


@register("cutpaper", "closing", alts=2)
def _cut_closing(k, slide, f, image):
    W, H, s, o = ctx(k)
    A = CUT_ART[k.ground]
    card = (0.22 * W, 0.16 * H, 0.56 * W, 0.36 * H) if o == "land" else (0.08 * W, 0.12 * H, 0.84 * W, 0.30 * H)
    if alt():                        # long copy: a roomier card (the paper sun sits behind it, as paper layers do)
        card = (0.10 * W, 0.08 * H, 0.80 * W, 0.56 * H) if o == "land" else (0.06 * W, 0.07 * H, 0.88 * W, 0.46 * H)
    pad = 0.4 * s
    items = [(x_, text_of(f, x_)) for x_ in ("title", "line") if text_of(f, x_)]
    r, d = _card_flow(k, slide, "closing", (card[0] + pad, card[1] + pad, card[2] - 2 * pad, card[3] - 2 * pad), items,
                      anchor="middle", align="c")
    na.paper_sun(slide, 0.85 * W, (0.14 if o == "land" else 0.50) * H, (0.16 * min(W, H), 0.11 * min(W, H)), A["sun"])
    _back_hills(k, slide, o, seed=21)
    na.paper_card(slide, *card)
    _front_hills(k, slide, o, seed=23)
    d()
    return r


@register("cutpaper", "points", alts=2)
def _cut_points(k, slide, f, image):
    import icons as _ic
    import tempfile as _tf
    W, H, s, o = ctx(k)
    A = CUT_ART[k.ground]
    pts = points_of(f.get("items"))
    icons_ = f.get("icons")
    if icons_ is not None and (not isinstance(icons_, (list, tuple)) or len(icons_) != len(pts)):
        raise ValueError("cutpaper.points(): icons= takes one icon spec per point ({} points), got {!r}".format(len(pts), icons_))
    rects, draws = {}, []
    head_items = [(x_, text_of(f, x_)) for x_ in ("kicker", "title") if text_of(f, x_)]
    r, d = flow(k, slide, "points", (0.07 * W, 0.08 * H, 0.86 * W, (0.16 if alt() == 0 else 0.30) * H), head_items)
    rects.update(r); draws.append(d)
    n = len(pts)
    head_bottom = max((v[1] + v[3] for v in r.values()), default=0.08 * H)
    if o == "land":
        gap = 0.35 * s
        top = max(0.28 * H, head_bottom + 0.3 * s) if alt() == 0 else head_bottom + 0.25 * s
        ch = 0.76 * H - top if alt() == 0 else 0.88 * H - top
        cw = (0.86 * W - (n - 1) * gap) / n
        cards = [(0.07 * W + i * (cw + gap), top, cw, ch) for i in range(n)]
    else:
        gap = 0.25 * s
        top = max(0.24 * H, head_bottom + 0.3 * s) if alt() == 0 else head_bottom + 0.25 * s
        cols = 1 if (alt() == 0 or n < 3) else 2              # long copy, 3-4 points: a 2-column grid of cards
        rows = -(-n // cols)
        ch = ((0.88 * H - top) - (rows - 1) * gap) / rows
        cwid = (0.84 * W - (cols - 1) * gap) / cols
        cards = [(0.08 * W + (i % cols) * (cwid + gap), top + (i // cols) * (ch + gap), cwid, ch) for i in range(n)]
    _front_hills(k, slide, o, seed=31, low=True)
    for i, ((cx, cy, cw, chh), (head, line)) in enumerate(zip(cards, pts)):
        na.paper_card(slide, cx, cy, cw, chh)
        compact = alt() > 0 and o != "land"
        dd = min(0.75 * s, chh * 0.30, cw * 0.24) if compact else min(1.1 * s, chh * 0.38, cw * 0.38)
        dx, dy = (cx + 0.25 * s + dd / 2, cy + 0.25 * s + dd / 2) if compact else (cx + 0.3 * s + dd / 2, cy + 0.3 * s + dd / 2)
        na.disc(slide, dx, dy, dd, A["discs"][i % len(A["discs"])], shadow=True)
        if icons_:
            png = str(Path(_tf.gettempdir()) / "slide-maker-native-art" / "icon_{}.png".format(str(icons_[i]).replace(":", "_")))
            Path(png).parent.mkdir(parents=True, exist_ok=True)
            _ic.icon_png(icons_[i], png, color="FFFFFF", px=200)
            draws.append(lambda png=png, dx=dx, dy=dy, dd=dd, head=head: dk.icon(slide, png, dx - dd * 0.3, dy - dd * 0.3, dd * 0.6, alt=head))
        else:
            nr, nd = flow(k, slide, "points", (dx - dd / 2, dy - dd / 2, dd, dd), [("mark", str(i + 1))], align="c",
                          anchor="middle", ink="FFFFFF", accent="FFFFFF", start={"mark": dd * 72 * 0.5})
            draws.append(nd)
        if compact:
            tcol = (cx + 0.45 * s + dd, cy + 0.2 * s, cw - dd - 0.65 * s, chh - 0.4 * s)
        elif o == "land":
            tcol = (cx + 0.3 * s, cy + 0.45 * s + dd, cw - 0.6 * s, chh - dd - 0.7 * s)
        else:
            tcol = (cx + 0.6 * s + dd, cy + 0.25 * s, cw - dd - 0.9 * s, chh - 0.5 * s)
        tr, td_ = _card_flow(k, slide, "points", tcol, [(x_, t) for x_, t in (("item_head", head), ("item_line", line)) if t],
                             anchor="top")
        draws.append(td_)
    _run_all(draws)
    return rects


# ═══════════════════════════════════ drafting 蓝图技术线稿 ═══════════════════════════════════
def _bp_ground(k, slide):
    na.grid_background(slide, base=k.P["ground"], ink=k.P["ink"])
    content, block = na.drawing_sheet(slide, ink=k.P["ink"], mute=k.P["mute"], accent=k.P["text_accents"][0],
                                      number=len(k.prs.slides), project=None, face=k.face("mono"))
    drawn = None
    if k.project and na.title_block_project(slide, block, k.project, ink=k.P["ink"], mute=k.P["mute"],
                                            face=k.face("mono")):
        drawn = k.project
    k._sheet = (content, block, drawn)                # the words this sheet's title block carries (None: none yet)


GROUNDS["drafting"] = _bp_ground


def _bp_area(k):
    """The sheet's content rect above the title block, and the block itself."""
    W, H, s, o = ctx(k)
    (cx, cy, cw, ch), block, _done = k._sheet
    return (cx, cy, cw, block[1] - 0.15 * s - cy), block


def _bp_project(k, slide, f, fallback=None):
    """Remember the caller's project words (project=, else the cover title) and make THIS sheet's title block carry
    them: written when the ground had none yet (the cover is where the title first arrives), rewritten in place when
    the caller names a new project on a later sheet."""
    explicit = text_of(f, "project")
    words = explicit or k.project or fallback
    if not words:
        return
    content, block, drawn = k._sheet
    if drawn == words:
        k.project = words
        return
    old = None
    if drawn is not None:                     # a new project on a later sheet: find the old words…
        for sh in list(slide.shapes):
            ps = sh.text_frame.paragraphs if getattr(sh, "has_text_frame", False) else []
            if len(ps) >= 2 and ps[0].text == "PROJECT" and ps[1].text == drawn:
                old = sh
                break
    if na.title_block_project(slide, block, words, ink=k.P["ink"], mute=k.P["mute"], face=k.face("mono")):
        if old is not None:                   # …and remove them only once the new ones are in (validate, then mutate)
            old._element.getparent().remove(old._element)
        k.project = words
        k._sheet = (content, block, words)
    elif explicit:
        raise vl.VLTextOverflow("drafting: project={!r} does not fit the sheet's title block even at 7pt — shorten "
                                "it".format(explicit[:40]))
    else:                                     # a remembered title too long for the block: the sheet number alone
        k.project = None
        k._sheet = (content, block, None)


@register("drafting", "cover")
def _bp_cover(k, slide, f, image):
    W, H, s, o = ctx(k)
    (cx, cy, cw, ch), _b = _bp_area(k)
    _bp_project(k, slide, f, fallback=text_of(f, "title"))
    col = (cx, cy + 0.10 * ch, cw * (0.62 if o == "land" else 1.0), ch * 0.80)
    items = [(x_, text_of(f, x_)) for x_ in ("kicker", "title", "subtitle") if text_of(f, x_)]
    r, d = flow(k, slide, "cover", col, items, anchor="middle", start={"title": 70 * s})
    d()
    return r


@register("drafting", "section")
def _bp_section(k, slide, f, image):
    W, H, s, o = ctx(k)
    (cx, cy, cw, ch), _b = _bp_area(k)
    _bp_project(k, slide, f)
    num, kicker, title = text_of(f, "number"), text_of(f, "kicker"), text_of(f, "title")
    rects, draws = {}, []
    dd = min(1.8 * s, ch * 0.5)
    by_ = cy + (ch - dd) / 2 if o == "land" else cy + 0.1 * ch
    if num:
        na.balloon(slide, cx, by_, dd, num, ink=k.P["ink"], accent=k.P["text_accents"][0], face=k.face("mono"),
                   fill=k.P["ground"])
    col = (cx + dd + 0.5 * s, cy, cw - dd - 0.5 * s, ch) if o == "land" else (cx, by_ + dd + 0.4 * s, cw, ch * 0.5)
    items = [(x_, t) for x_, t in (("kicker", kicker), ("title", title)) if t]
    r, d = flow(k, slide, "section", col, items, anchor="middle" if o == "land" else "top")
    rects.update(r); draws.append(d)
    _run_all(draws)
    return rects


@register("drafting", "image_text")
def _bp_image_text(k, slide, f, image):
    W, H, s, o = ctx(k)
    (cx, cy, cw, ch), _b = _bp_area(k)
    _bp_project(k, slide, f)
    img = (cx, cy, cw * 0.56, ch) if o == "land" else (cx, cy, cw, ch * 0.50)
    col = (cx + cw * 0.62, cy + 0.1 * ch, cw * 0.38, ch * 0.8) if o == "land" else (cx, cy + ch * 0.56, cw, ch * 0.44)
    items = [(x_, text_of(f, x_)) for x_ in ("kicker", "title", "body", "caption") if text_of(f, x_)]
    r, d = flow(k, slide, "image_text", col, items, anchor="middle" if o == "land" else "top")
    place_image(k, slide, image, img, "image_text")
    d()
    return r


@register("drafting", "quote")
def _bp_quote(k, slide, f, image):
    W, H, s, o = ctx(k)
    (cx, cy, cw, ch), _b = _bp_area(k)
    _bp_project(k, slide, f)
    has_img = image is not None
    col = (cx, cy + 0.12 * ch, cw * (0.58 if has_img and o == "land" else 0.85), ch * 0.72)
    items = [(x_, text_of(f, x_)) for x_ in ("quote", "attribution") if text_of(f, x_)]
    r, d = flow(k, slide, "quote", col, items, anchor="middle")
    if has_img:
        place_image(k, slide, image, (cx + cw * 0.64, cy, cw * 0.36, ch) if o == "land" else (cx, cy + ch * 0.6, cw, ch * 0.4),
                    "quote")
    d()
    return r


@register("drafting", "data")
def _bp_data(k, slide, f, image):
    W, H, s, o = ctx(k)
    (cx, cy, cw, ch), _b = _bp_area(k)
    _bp_project(k, slide, f)
    num, label, note = text_of(f, "number"), text_of(f, "label"), text_of(f, "note")
    rects, draws = {}, []
    nrect = (cx, cy, cw * 0.40, ch) if o == "land" else (cx, cy, cw, ch * 0.48)
    if num:
        r, d = flow(k, slide, "data", nrect, [("number", num)], anchor="middle", align="c")
        rects.update(r); draws.append(d)
        nr = r.get("number", nrect)
        # beside the numeral's INK (it is centred in its rect): the plan clamped x to 46% of the width, which put the
        # line through the "3" on a portrait sheet
        import display_type as _dt
        nsz = (nr[3] - 0.06) / 1.2 * 72.0                  # one line: _field_height's model, inverted
        ink_w = _dt._glyph_width(num, nsz, k.face("numeral"), vl.TYPE[k.name]["number"][2]) or nr[2] * 0.5
        dxl = nr[0] + nr[2] / 2.0 + ink_w / 2.0 + 0.3 * s
        draws.append(lambda: na.dimension_line(slide, dxl, nr[1] + 0.1 * nr[3], nr[1] + 0.9 * nr[3], ink=k.P["ink"]))
    col = (cx + cw * 0.50, cy + 0.15 * ch, cw * 0.50, ch * 0.7) if o == "land" else (cx, cy + ch * 0.54, cw, ch * 0.44)
    items = [(x_, t) for x_, t in (("label", label), ("note", note)) if t]
    r, d = flow(k, slide, "data", col, items, anchor="middle" if o == "land" else "top")
    rects.update(r); draws.append(d)
    _run_all(draws)
    return rects


@register("drafting", "closing")
def _bp_closing(k, slide, f, image):
    W, H, s, o = ctx(k)
    (cx, cy, cw, ch), _b = _bp_area(k)
    _bp_project(k, slide, f)
    col = (cx, cy + 0.15 * ch, cw * (0.70 if o == "land" else 1.0), ch * 0.6)
    items = [(x_, text_of(f, x_)) for x_ in ("title", "line") if text_of(f, x_)]
    r, d = flow(k, slide, "closing", col, items, anchor="middle")
    d()
    return r


@register("drafting", "points", alts=3)
def _bp_points(k, slide, f, image):
    """One iso layer per point, each with a numbered leader to its note. The notes take EVEN slots of the free
    height, not the layers' gap (measured 2026-10-05: tied to the gap, a note got 0.47in and ordinary copy was
    refused); the leaders angle from each layer's right corner to its balloon."""
    W, H, s, o = ctx(k)
    (cx, cy, cw, ch), _b = _bp_area(k)
    _bp_project(k, slide, f)
    pts = points_of(f.get("items"))
    n = len(pts)
    rects, draws = {}, []
    head_items = [(x_, text_of(f, x_)) for x_ in ("kicker", "title") if text_of(f, x_)]
    hw = cw * (0.62 if o == "land" and alt() == 0 else 1.0)
    hh = ch * ((0.30 if o == "land" else 0.20) if alt() == 0 else (0.36 if o == "land" else 0.26) if alt() == 1
               else (0.34 if o == "land" else 0.40))
    # the parts-legend layout is the last resort for long copy: its title starts near its floor and takes only the
    # height it needs, leaving the rest to the notes
    r, d = flow(k, slide, "points", (cx, cy, hw, hh), head_items, start={"title": 30 * s} if alt() == 2 else None)
    rects.update(r); draws.append(d)
    top = max((v[1] + v[3] for v in r.values()), default=cy) + (0.35 if alt() < 2 else 0.2) * s
    avail = cy + ch - top                                  # above the title block: where the notes go
    thick = 0.18
    if o == "land":                                        # the block is at the right: the drawing may go lower
        sheet = k._sheet[0]
        stack_avail = sheet[1] + sheet[3] - top
        size = min(cw * 0.26, stack_avail * 0.46)         # the drawing fills the height: the plates, then air between
        scx, notes_x = cx + cw * 0.26, cx + cw * 0.56
        if alt():                                          # long copy: a smaller drawing, the notes get the width
            size, scx, notes_x = min(cw * 0.19, stack_avail * 0.40), cx + cw * 0.19, cx + cw * 0.42
    else:
        stack_avail = avail
        size = min(cw * 0.30, avail * 0.34)
        scx, notes_x = cx + cw * 0.30, cx + cw * 0.62
        if alt() == 1:
            size, scx, notes_x = min(cw * 0.24, avail * 0.30), cx + cw * 0.24, cx + cw * 0.50
        elif alt() == 2:                                   # the parts legend needs the width more than the drawing
            size, scx, notes_x = min(cw * 0.13, avail * 0.30), cx + cw * 0.13, cx + cw * 0.30
    gap = min(0.9 * s, max(0.35 * s, (stack_avail - size - thick - 0.2 * s) / max(n - 1, 1)))
    tops = na.iso_stack(slide, scx, top, size, gap, n, ink=k.P["ink"], accent=k.P["accents"][0], accent_layer=0,
                        thick=thick)
    bd = 0.42 * s
    extent = size + (n - 1) * gap + thick                  # the notes span the stack they label (room for their words)
    slot = min(avail, max(extent, n * 1.05 * s)) / n
    tx = notes_x + bd + 0.2 * s
    if alt() == 2:                 # a drawing's parts list: a balloon on each layer, the notes in a numbered legend
        bd = 0.36 * s                                          # a parts list's balloons are smaller than leaders'
        bdl = min(bd, gap * 0.9)                               # layer balloons never touch their neighbours
        notes_x = max(notes_x, scx + size * 0.866 + 0.08 * s + bdl + 0.35 * s)   # nor the legend's balloons
        cols = 2 if (n >= 3 and (cx + cw - notes_x) / 2 >= 1.6 * s) else 1     # two columns only when each is wide
        rows = -(-n // cols)
        lw = (cx + cw - notes_x - (cols - 1) * 0.2 * s) / cols
        lh = avail / rows
        for i, ((head, line), layer) in enumerate(zip(pts, tops)):
            ax, ay = layer[1]
            lx, ly = notes_x + (i % cols) * (lw + 0.2 * s), top + (i // cols) * lh
            tr, td_ = flow(k, slide, "points", (lx + bd + 0.15 * s, ly - 0.05 * s, lw - bd - 0.15 * s, lh - 0.03 * s),
                           [(x_, t) for x_, t in (("item_head", head), ("item_line", line)) if t], anchor="top")
            for bx_, by2, dd_ in ((ax + 0.08 * s, ay - bdl / 2, bdl), (lx, ly, bd)):
                draws.append(lambda bx_=bx_, by2=by2, dd_=dd_, i=i: na.balloon(
                    slide, bx_, by2, dd_, str(i + 1), ink=k.P["ink"], accent=k.P["text_accents"][0],
                    face=k.face("mono"), fill=k.P["ground"]))
            draws.append(td_)
        _run_all(draws)
        return rects
    for i, ((head, line), layer) in enumerate(zip(pts, tops)):
        ax, ay = layer[1]                                     # the layer's right corner
        by_ = top + i * slot
        tr, td_ = flow(k, slide, "points", (tx, by_ - 0.05 * s, cx + cw - tx, slot * 0.92),
                       [(x_, t) for x_, t in (("item_head", head), ("item_line", line)) if t], anchor="top")
        draws.append(lambda ax=ax, ay=ay, by_=by_: na.seg(slide, ax, ay, notes_x, by_ + bd / 2, k.P["ink"], w=0.6))
        draws.append(lambda by_=by_, i=i: na.balloon(slide, notes_x, by_, bd, str(i + 1), ink=k.P["ink"],
                                                     accent=k.P["text_accents"][0], face=k.face("mono"),
                                                     fill=k.P["ground"]))
        draws.append(td_)
    _run_all(draws)
    return rects
