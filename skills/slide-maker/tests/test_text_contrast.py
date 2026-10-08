#!/usr/bin/env python3
"""WCAG 1.4.3 text contrast as a hand-off floor: text under 14pt (or under 18pt and not bold) needs 4.5:1 on a resolved backing — no judgement needed, so the a11y gate holds it (a low-contrast caption passed --gate-check clean)."""
from __future__ import annotations
import sys, tempfile
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
fails: list[str] = []


def check(cond, msg):
    if not cond:
        fails.append(msg)



import subprocess
import deckkit as dk, lint_deck as ld
td = Path(tempfile.mkdtemp())
MUTE, CREAM = dk.RGBColor.from_string("8A7A63"), "F5E6C8"          # 3.38:1 — the P1 run's caption
prs = dk.blank_deck(13.333, 7.5)
for size, bold, label in ((12.5, False, "small caption"), (15, False, "fifteen point regular"), (15, True, "fifteen point bold"),
                          (20, False, "twenty point regular")):
    s = dk.add_slide(prs)
    dk.box(s, 0.5, 0.5, 12.3, 6.5, fill=CREAM)
    dk.text(s, 1, 1, 10, 1.2, [[("A " + label + " in the muted ink on cream", size, MUTE, bold, False)]])
deck = td / "deck.pptx"
prs.save(str(deck))
r = subprocess.run([sys.executable, str(ROOT / "scripts" / "lint_deck.py"), str(deck), "--static"], capture_output=True, text=True)
out = r.stdout + r.stderr
import re
def fired(n):                                   # not "NON-TEXT CONTRAST" (a substring of it)
    return [l_ for l_ in out.splitlines() if re.search(r"(?<!NON-)TEXT CONTRAST:", l_)
            and ("slide {}:".format(n) in l_ or "slide  {}".format(n) in l_)]
check(fired(1), "12.5pt text at 3.38:1 fails WCAG 1.4.3 (under 14pt is never large text): {}".format(out[-400:]))
check(fired(2), "15pt regular text at 3.38:1 fails (large needs 18pt, or 14pt BOLD)")
check(not fired(3), "15pt BOLD is large text — 3:1 suffices: {}".format(fired(3)))
# the finding names the nearest ink that keeps the hue and clears the floor — a remedy any agent can paste
_sug = re.findall(r"#([0-9A-F]{6}) keeps", (fired(1) or [""])[0])
check(_sug and dk.contrast_ratio(dk.RGBColor.from_string(_sug[0]), CREAM) >= 4.5,
      "TEXT CONTRAST suggests an ink that reaches 4.5:1 on the fill: {}".format((fired(1) or [""])[0][-160:]))
check(not fired(4), "20pt is large text — 3:1 suffices: {}".format(fired(4)))
check("TEXT CONTRAST" in ld.A11Y_WCAG and "TEXT CONTRAST" in ld.A11Y_BLOCKING and "TEXT CONTRAST" in ld.A11Y_CODES,
      "the hand-off a11y floors hold it on both runtimes (it is in A11Y_WCAG / A11Y_BLOCKING)")

# A size or weight that is INHERITED was not read: the run is not judged as 12pt (the lint's fallback) — a template
# title or raw python-pptx text at the deck's 18pt default is large text and clears WCAG at 3:1 (final review: a
# 3.45:1 default-size headline was held as "12pt text"). Unknown -> the advisory band, never a blocking guess.
from pptx.util import Inches as _In, Pt as _Pt
from pptx.enum.shapes import MSO_SHAPE as _MS
prs = dk.blank_deck(13.333, 7.5)
s_ = dk.add_slide(prs)
sh_ = s_.shapes.add_shape(_MS.RECTANGLE, _In(1), _In(2), _In(6), _In(1.5))
sh_.fill.solid(); sh_.fill.fore_color.rgb = dk.RGBColor.from_string("FFFFFF")
r_ = sh_.text_frame.paragraphs[0].add_run(); r_.text = "Large default-size headline"
r_.font.color.rgb = dk.RGBColor.from_string("8A8A8A")                 # 3.45:1, no size, no bold of its own
s_ = dk.add_slide(prs)
sh_ = s_.shapes.add_shape(_MS.RECTANGLE, _In(1), _In(2), _In(6), _In(1.5))
sh_.fill.solid(); sh_.fill.fore_color.rgb = dk.RGBColor.from_string("FFFFFF")
r_ = sh_.text_frame.paragraphs[0].add_run(); r_.text = "Fifteen point, weight inherited"
r_.font.size = _Pt(15); r_.font.color.rgb = dk.RGBColor.from_string("8A8A8A")
s_ = dk.add_slide(prs)
sh_ = s_.shapes.add_shape(_MS.RECTANGLE, _In(1), _In(2), _In(6), _In(1.5))
sh_.fill.solid(); sh_.fill.fore_color.rgb = dk.RGBColor.from_string("FFFFFF")
r_ = sh_.text_frame.paragraphs[0].add_run(); r_.text = "Twelve point, weight inherited"
r_.font.size = _Pt(12); r_.font.color.rgb = dk.RGBColor.from_string("8A8A8A")
inh = td / "inherited.pptx"
prs.save(str(inh))
r = subprocess.run([sys.executable, str(ROOT / "scripts" / "lint_deck.py"), str(inh), "--static"], capture_output=True, text=True)
out = r.stdout + r.stderr
check(not fired(1), "a run with NO size of its own is not judged as 12pt: {}".format(fired(1)))
check(not fired(2), "15pt with an inherited weight may be bold (large) — advisory, not held: {}".format(fired(2)))
check(fired(3), "12pt is normal text at any weight — still held: {}".format(out[-300:]))
check(not any("12pt" in l_ for l_ in fired(1) + fired(2)), "no size is printed that was not read")

# deckkit's OWN helpers must clear the floor with their defaults — the callout's 11pt bold label
# was MAGENTA on TINT at 4.27:1, so every deck with a takeaway bar would now be held at hand-off.
# The label text reaches 4.5:1; the accent BAR keeps the exact accent (it is non-text, 3:1).
def _lint_static(p):
    r_ = subprocess.run([sys.executable, str(ROOT / "scripts" / "lint_deck.py"), str(p), "--static"],
                        capture_output=True, text=True)
    return [l_ for l_ in (r_.stdout + r_.stderr).splitlines() if re.search(r"(?<!NON-)TEXT CONTRAST:", l_)]

def _runs_and_bars(slide):
    runs, fills = [], []
    for sh in slide.shapes:
        if sh.has_text_frame:
            runs += [r_ for p_ in sh.text_frame.paragraphs for r_ in p_.runs]
        try:
            if sh.fill.type == 1:
                fills.append(str(sh.fill.fore_color.rgb))
        except Exception:
            pass
    return runs, fills

GOLD_ON_TINT = (dk.GOLD, dk.TINT)                       # a user accent far under 4.5:1 on the tint
TEAL_ON_DEEP = (dk.TEAL, dk.DEEP)                       # a light accent on a DARK card: must LIGHTEN
prs = dk.blank_deck()
s1 = dk.add_slide(prs); dk.callout(s1, 0.5, 1.0, 9.0, 0.6, "TAKEAWAY", "the default callout")
s2 = dk.add_slide(prs); dk.bottom_callout(s2, 0.5, 9.0, "WHY", "the footer-safe callout")
s3 = dk.add_slide(prs); dk.callout(s3, 0.5, 1.0, 9.0, 0.6, "NOTE", "a gold accent", label_c=GOLD_ON_TINT[0])
s4 = dk.add_slide(prs); dk.callout(s4, 0.5, 1.0, 9.0, 0.6, "NOTE", "on a dark card", label_c=TEAL_ON_DEEP[0],
                                   fill=TEAL_ON_DEEP[1], body_c=dk.WHITE)
s5 = dk.add_slide(prs); dk.callout(s5, 0.5, 1.0, 9.0, 0.6, "NOTE", "already legible", label_c=dk.DEEP)
s6 = dk.add_slide(prs)
dk.consort_flow(s6, 0.5, 0.4, 9.0, 4.6, [("Screened", 200, [("not eligible", 14)]),
                                          ("Enrolled", 186, [("withdrew", 52)]), ("Analysed", 134)])
hd = td / "helpers.pptx"
prs.save(str(hd))
hits = _lint_static(hd)
check(not hits, "deckkit's own callout / bottom_callout / consort_flow defaults clear WCAG 1.4.3: {}".format(hits))
from pptx import Presentation as _P
_sl = _P(str(hd)).slides
for i_, (acc_, fill_) in ((0, (dk.MAGENTA, dk.TINT)), (2, GOLD_ON_TINT), (3, TEAL_ON_DEEP)):
    runs_, fills_ = _runs_and_bars(_sl[i_])
    lab = runs_[0]
    ratio = dk.contrast_ratio(lab.font.color.rgb, fill_)
    check(ratio >= 4.5, "slide {}: the label text reaches 4.5:1 on its card (got {:.2f})".format(i_ + 1, ratio))
    check(str(acc_) in fills_, "slide {}: the accent bar keeps the exact accent #{} (bars {})".format(i_ + 1, acc_, fills_))
    check(str(lab.font.color.rgb) != str(acc_), "slide {}: the label was adjusted, not left at the accent".format(i_ + 1))
_lab4 = _runs_and_bars(_sl[3])[0][0].font.color.rgb
check(sum(_lab4) > sum(dk.TEAL), "on a DARK card the label moves toward white, not black (got #{})".format(_lab4))
check(str(_runs_and_bars(_sl[4])[0][0].font.color.rgb) == str(dk.DEEP),
      "an accent that already clears 4.5:1 is written unchanged (byte-identical output)")
for c_, bg_ in ((dk.MAGENTA, dk.TINT), (dk.GOLD, dk.WHITE), (dk.TEAL, dk.DEEP), (dk.PALE, dk.DEEP),
                (dk.RGBColor.from_string("808080"), dk.RGBColor.from_string("808080"))):
    if not hasattr(dk, "_ink_reaching"):
        check(False, "deckkit._ink_reaching exists"); break
    got = dk._ink_reaching(c_, bg_, 4.5)
    check(dk.contrast_ratio(got, bg_) >= 4.5, "_ink_reaching #{} on #{} reaches 4.5:1 (got {:.2f})".format(
        c_, bg_, dk.contrast_ratio(got, bg_)))

# callout(fill=None) — an unfilled callout — built before the label ink learned its fill; it must still build, the
# label in the caller's own accent (there is no fill to measure against)
try:
    _p0 = dk.blank_deck(); _s0 = dk.add_slide(_p0)
    dk.callout(_s0, 0.5, 1.0, 9.0, 0.6, "NOTE", "an unfilled callout", fill=None)
    _lab0 = [r_ for sh in _s0.shapes if sh.has_text_frame for p_ in sh.text_frame.paragraphs for r_ in p_.runs][0]
    check(str(_lab0.font.color.rgb) == str(dk.MAGENTA), "callout(fill=None) keeps the accent label: #{}".format(_lab0.font.color.rgb))
except Exception as e:
    check(False, "callout(fill=None) builds (it did before this branch): {!r}".format(e))

# Both runtimes say the RIGHT remedy per code: TEXT CONTRAST is 4.5:1 for text (WCAG 1.4.3), not "3:1 … why this mark
# is decorative" — a non-Claude agent following that text aims for the wrong floor (final review 2026-10-04)
check(set(ld.A11Y_BLOCKING) <= set(getattr(ld, "A11Y_REMEDY", {})), "every blocking a11y code has a remedy line")
check("4.5:1" in getattr(ld, "A11Y_REMEDY", {}).get("TEXT CONTRAST", ""), "the TEXT CONTRAST remedy names 4.5:1")
for _rt in ("codex_delivery_gate.py", "render_deck.py"):
    _src = (ROOT / "scripts" / _rt).read_text(encoding="utf-8")
    check("A11Y_REMEDY" in _src, "{} prints the per-code remedy (lint_deck.A11Y_REMEDY)".format(_rt))
check("WCAG 1.4.11 3:1 floor" not in (ROOT / "scripts" / "codex_delivery_gate.py").read_text(encoding="utf-8"),
      "the codex gate no longer calls every a11y code a 3:1 mark floor")

# Property: on ANY ground, the inks deckkit picks reach 4.5:1 — black or white always does (>= 4.58:1), so "no ink
# reaches" is never true; a near-black #111111 pole printed it for #1F8A70 / #E0392B (final review).
_grid = [dk.RGBColor(r_, g_, b_) for r_ in range(0, 256, 17) for g_ in range(0, 256, 17) for b_ in range(0, 256, 17)]
_bad_leg = [str(g_) for g_ in _grid if dk.contrast_ratio(dk._legible_ink(g_), g_) < 4.5]
check(not _bad_leg, "_legible_ink reaches 4.5:1 on every ground of a 16^3 grid: {} fail, e.g. {}".format(len(_bad_leg), _bad_leg[:4]))
_bad_reach = []
for g_ in _grid[::7]:
    for c_ in (dk.MAGENTA, dk.WHITE, dk.GOLD, dk.RGBColor(0x1F, 0x8A, 0x70)):
        try:
            if dk.contrast_ratio(dk._ink_reaching(c_, g_, 4.5), g_) < 4.5:
                _bad_reach.append((str(c_), str(g_)))
        except ValueError:
            _bad_reach.append((str(c_), str(g_), "raised"))
check(not _bad_reach, "_ink_reaching reaches 4.5:1 on every ground (never raises): {}".format(_bad_reach[:4]))

# Every shipped preset: the helpers that set small text on their own fills clear WCAG 1.4.3 (the final review found
# consort_flow, corner_tab (via eval_matrix), org_tree and segmented_bar held under editorial_paper / consulting / eastern_traditional
# / luxury_dark / editorial_report / memphis — the default-palette sweep could not see it).
import presets, sigs, io as _io, contextlib as _cl
_snap = dk._state_snapshot()
_held = []
for _pn in presets.PRESETS:
    dk._state_restore(_snap)
    with _cl.redirect_stdout(_io.StringIO()), _cl.redirect_stderr(_io.StringIO()):
        presets.apply(_pn)
    _prs = dk.blank_deck(10, 5.625)
    for _ex in ("consort_flow", "eval_matrix", "org_tree", "segmented_bar"):   # eval_matrix carries corner_tab
        _sl = _prs.slides.add_slide(_prs.slide_layouts[6])
        with _cl.redirect_stdout(_io.StringIO()):
            exec(compile(sigs.EXAMPLES[_ex], _ex, "exec"), {"dk": dk, "s": _sl, "prs": _prs})
    _sl = _prs.slides.add_slide(_prs.slide_layouts[6]); dk.callout(_sl, 0.5, 1.0, 9.0, 0.6, "NOTE", "a callout")
    _f = td / "preset-{}.pptx".format(_pn)
    _prs.save(str(_f))
    _held += ["{}: {}".format(_pn, l_.strip()[:120]) for l_ in _lint_static(_f)]
dk._state_restore(_snap)
check(not _held, "every preset: deckkit's helpers raise no TEXT CONTRAST ({}): {}".format(len(_held), _held[:4]))

print("\n".join("FAIL " + f for f in fails) if fails else "", end="")
print("[test_text_contrast] {}".format("FAILED: {} problem(s)".format(len(fails)) if fails else "ok"))
sys.exit(1 if fails else 0)
