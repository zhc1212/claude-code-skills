#!/usr/bin/env python3
"""deckkit.a11y_title(slide, text): the slide's title for screen readers — a TITLE placeholder, first in reading order,
off the canvas so nothing is drawn. The docs called an off-canvas title "the sanctioned trick for statement slides" and
no helper made one: a docs-only agent's text box tripped OFF_CANVAS + DUPLICATE_TEXT and blocked the build, and every
visual-language page that set its kicker before its title (or its title below the top 28%) was held at hand-off on
READING ORDER or reported NO SLIDE TITLE (restricted non-Claude run, 2026-10-04)."""
from __future__ import annotations
import io, contextlib, re, sys, tempfile
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
fails: list[str] = []


def check(cond, msg):
    if not cond:
        fails.append(msg)


import deckkit as dk
import lint_deck as ld
import visual_languages as vl
from pptx.enum.shapes import PP_PLACEHOLDER

td = Path(tempfile.mkdtemp())


def lint_lines(path):
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        ld.lint(str(path), mode="presented", static_ok=True)
    return buf.getvalue().splitlines()


def code_on(lines, code, n):
    return [l_ for l_ in lines if re.search(r"slide\s+{}:.*\]\s+{}:".format(n, re.escape(code)), l_)]


# 1. the helper: a TITLE placeholder, first in the shape tree, wholly off the canvas, carrying the text
prs = dk.blank_deck(13.333, 7.5)
s = dk.add_slide(prs)
dk.add_slide(prs)                                   # slide 1 is exempt from title checks; test slide 2
s = prs.slides[1]
dk.text(s, 1.0, 3.0, 8.0, 0.5, [[("A KICKER", 12, dk.MAGENTA, True, False)]])
dk.text(s, 1.0, 3.6, 10.0, 1.4, [[("Bring it broken, take it home working", 40, dk.DEEP, True, False)]])
sh = dk.a11y_title(s, "Bring it broken, take it home working")
check(sh.is_placeholder and sh.placeholder_format.type == PP_PLACEHOLDER.TITLE, "a11y_title is a TITLE placeholder")
check(list(s.shapes)[0].shape_id == sh.shape_id, "a11y_title is FIRST in reading order")
H = prs.slide_height
check(sh.top + sh.height <= 0 or sh.left >= prs.slide_width or sh.left + sh.width <= 0 or sh.top >= H,
      "a11y_title sits wholly off the canvas (nothing is drawn): {} {} {} {}".format(sh.left, sh.top, sh.width, sh.height))
check(sh.text_frame.text == "Bring it broken, take it home working", "a11y_title carries the text")
bad = [f for f in dk.lint_layout(prs, verbose=False) if f[2] in ("OFF_CANVAS", "DUPLICATE_TEXT") or f[1] == "CRITICAL"]
check(not bad, "lint_layout does not flag the declared title (OFF_CANVAS / DUPLICATE_TEXT): {}".format(bad))
p = td / "t.pptx"
prs.save(str(p))
L = lint_lines(p)
check(not code_on(L, "NO SLIDE TITLE", 2) and not code_on(L, "READING ORDER", 2),
      "lint_deck reads it as the title, first in order: {}".format([l_ for l_ in L if "slide 2" in l_][:4]))
try:
    dk.a11y_title(s, "   ")
    fails.append("an empty a11y_title was accepted")
except ValueError:
    pass

# 2. every visual-language page declares its title — the agent's own long Chinese copy, every language, both grounds
CONTRAST = {"editorial": "ink", "soft": "dusk", "collage": "slate", "storybook": "meadow"}
ph = str(ROOT / "assets" / "vl" / "photo" / "hall-repair.jpg")
wc = str(ROOT / "assets" / "vl" / "watercolour" / "rooftop-garden.jpg")
for name, var in CONTRAST.items():
    for g in ("light", var):
        prs = dk.blank_deck(13.333, 7.5)
        with contextlib.redirect_stdout(io.StringIO()):
            k = vl.use(name, prs, ground=g)
        img = wc if name == "storybook" else ph
        k.cover(k.new_slide(), kicker="社区修理咖啡馆", title="坏了别扔，带来一起修", image=[ph, ph] if name == "collage" else img)
        k.section(k.new_slide(), number="01", kicker="这是什么", title="一个让东西重新用起来的晚上")
        k.image_text(k.new_slide(), kicker="一个晚上如何进行", title="带上东西，找一张桌子，志愿者在旁边带你一起修",
                     body="没有预约，也不收费。", image=img)
        k.image_text(k.new_slide(), kicker="大家带来什么", title="小家电、衣服、玩具、家具——能在桌上修的，都可以带来",
                     body="不设限。", image=img)
        k.quote(k.new_slide(), quote="工具在来访者手里，志愿者只是在旁边指点。", attribution="一位志愿者")
        k.data(k.new_slide(), number="1", label="每月一个晚上", note="够短，也够常。")
        k.closing(k.new_slide(), title="下个月，带一样坏东西来", line="也带上一位邻居。")
        p = td / "vl-{}-{}.pptx".format(name, g)
        prs.save(str(p))
        L = lint_lines(p)
        held = [l_.strip()[:110] for l_ in L if re.search(r"\]\s+(NO SLIDE TITLE|READING ORDER|DUPLICATE SLIDE TITLES):", l_)]
        check(not held, "{} {}: every kit page has its title, first in order: {}".format(name, g, held[:3]))
        bad = [f for f in dk.lint_layout(prs, verbose=False) if f[2] in ("OFF_CANVAS", "DUPLICATE_TEXT")]
        check(not bad, "{} {}: the declared titles raise no OFF_CANVAS / DUPLICATE_TEXT: {}".format(name, g, bad[:2]))

print("\n".join("FAIL " + f for f in fails) if fails else "", end="")
print("[test_a11y_title] {}".format("FAILED: {} problem(s)".format(len(fails)) if fails else "ok"))
sys.exit(1 if fails else 0)
