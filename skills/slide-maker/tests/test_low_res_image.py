#!/usr/bin/env python3
"""LOW_RES_IMAGE: a picture whose visible source pixels are spread too thin over its frame (< 72 ppi WARN, < 36 ppi CRITICAL; long side >= 1.5in; decorative and declared pixel art exempt)."""
from __future__ import annotations
import sys, tempfile
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
fails: list[str] = []


def check(cond, msg):
    if not cond:
        fails.append(msg)



import deckkit as dk
from PIL import Image
td = Path(tempfile.mkdtemp())


def img(w, h, name):
    p = td / name
    im = Image.new("RGB", (w, h))
    im.putdata([(60 + (x * 7 + y * 3) % 150, 90 + (x * 5) % 120, 40 + (y * 9) % 160) for y in range(h) for x in range(w)])
    im.save(p)
    return str(p)


def lint_codes(prs):
    return [(f[0], f[1], f[2], f[3]) for f in dk.lint_layout(prs, verbose=False) if f[2] == "LOW_RES_IMAGE"]


prs = dk.blank_deck(13.333, 7.5)
s = dk.add_slide(prs); dk.picture(s, img(120, 80, "a.png"), 1, 1, 6, 4, alt="a tiny source on a big frame")      # 20 ppi
s = dk.add_slide(prs); dk.picture(s, img(300, 200, "b.png"), 1, 1, 6, 4, alt="a web thumbnail")                # 50 ppi
s = dk.add_slide(prs); dk.picture(s, img(1200, 800, "c.png"), 1, 1, 6, 4, alt="a proper photo")               # 200 ppi
s = dk.add_slide(prs); dk.picture(s, img(30, 20, "d.png"), 1, 1, 1.2, 0.8, alt="a small thumbnail")          # small: not checked
s = dk.add_slide(prs)
p = dk.picture(s, img(1200, 800, "e.png"), 1, 1, 6, 4, alt="a big photo cropped to a sliver")
p.crop_left, p.crop_right = 0.45, 0.45                                                                      # 120 px visible → 20 ppi
s = dk.add_slide(prs); dk.decorative(dk.picture(s, img(64, 64, "f.png"), 0, 0, 6, 6, alt="paper"), "a paper texture under the page; nothing rides on its detail")
s = dk.add_slide(prs); dk.low_res_intent(dk.picture(s, img(32, 32, "g.png"), 1, 1, 4, 4, alt="pixel art"), "the pixel-art sprite is the subject; the blocks are the point")
f = lint_codes(prs)
by = {n: (sev, msg) for n, sev, _c, msg in f}
check(by.get(1, ("",))[0] == "CRITICAL", "20 ppi is CRITICAL: {}".format(by.get(1)))
check(by.get(2, ("",))[0] == "WARN", "50 ppi is a WARN: {}".format(by.get(2)))
check(3 not in by and 4 not in by, "200 ppi and a small thumbnail are not reported: {}".format(f))
check(by.get(5, ("",))[0] == "CRITICAL", "a crop that leaves 120 px visible is CRITICAL: {}".format(by.get(5)))
check(6 not in by and 7 not in by, "decorative and declared pixel art are exempt: {}".format(f))
m = by.get(1, ("", ""))[1]
check("120" in m and "ppi" in m and "low_res_intent" in m, "the message gives the pixels, the ppi and the remedies: {}".format(m))
try:
    dk.low_res_intent(dk.picture(dk.add_slide(prs), img(32, 32, "h.png"), 1, 1, 4, 4, alt="x"), "ok")
    fails.append("low_res_intent accepted a reason nobody can disagree with")
except ValueError:
    pass
# the declaration composes with the others (a generated series picture stays a series picture)
q = dk.picture(dk.add_slide(prs), img(32, 32, "i.png"), 1, 1, 4, 4, alt="y")
dk.overlap_intent(q, "rides under the headline on purpose")
dk.low_res_intent(q, "the pixel-art sprite is the subject; the blocks are the point")
check("+overlap" in q.name or q.name.startswith("deckkit-overlap"), "overlap declaration kept: {}".format(q.name))
check("lowres" in q.name, "low-res declaration added: {}".format(q.name))
# strict lint refuses a CRITICAL one
prs2 = dk.blank_deck(13.333, 7.5)
dk.picture(dk.add_slide(prs2), img(120, 80, "j.png"), 1, 1, 6, 4, alt="tiny")
try:
    dk.lint_layout(prs2, verbose=False, strict=True)
    fails.append("strict lint saved a 20 ppi picture")
except Exception:
    pass

# A picture PLACEHOLDER filled with the user's photo (the template path, where low-res sources are likeliest) and a
# picture inside a SCALED group are measured at the size they are drawn (final review 2026-10-04: both were silent).
from PIL import Image as _Im
from pptx import Presentation as _Pr
from pptx.util import Inches as _In
_Im.new("RGB", (400, 300), (200, 80, 60)).save(str(td / "small400.png"))
_pp = _Pr()
_pp.slide_width, _pp.slide_height = _In(13.333), _In(7.5)
_ps = _pp.slides.add_slide(_pp.slide_layouts[8])                 # "Picture with Caption": a picture placeholder
_ph = [q for q in _ps.placeholders if q.placeholder_format.type == 18][0]
_ph.width, _ph.height = _In(6), _In(4.5)
_ph.insert_picture(str(td / "small400.png"))
_f = dk._low_res_findings(_ps, 1)
check(_f and _f[0][1] == "WARN" and "67 ppi" in _f[0][3], "a 400x300 photo in a 6x4.5in picture placeholder (67 ppi) warns: {}".format(_f))
_gp = dk.blank_deck(13.333, 7.5)
_gs = dk.add_slide(_gp)
_g = _gs.shapes.add_group_shape()
_g.shapes.add_picture(str(td / "small400.png"), _In(0), _In(0), _In(4), _In(3))
_gx = _g._element.grpSpPr.find(dk.qn("a:xfrm"))
_gx.ext.cx, _gx.ext.cy = _In(10), _In(7.5)                       # drawn at 10x7.5in ...
_gx.chExt.cx, _gx.chExt.cy = _In(4), _In(3)                      # ... from a 4x3in child space
_f = dk._low_res_findings(_gs, 1)
check(_f and "40 ppi" in _f[0][3] and "10.0x7.5in" in _f[0][3], "a picture in a group scaled 2.5x is measured as drawn (40 ppi): {}".format(_f))
_gs2 = dk.add_slide(_gp)
_g2 = _gs2.shapes.add_group_shape()
_g2.shapes.add_picture(str(td / "small400.png"), _In(0), _In(0), _In(10), _In(7.5))
_gx2 = _g2._element.grpSpPr.find(dk.qn("a:xfrm"))
_gx2.ext.cx, _gx2.ext.cy = _In(4), _In(3)                        # a 10in child drawn at 4in: 100 ppi, fine
_gx2.chExt.cx, _gx2.chExt.cy = _In(10), _In(7.5)
check(not dk._low_res_findings(_gs2, 2), "a group scaled DOWN is not flagged at its child size: {}".format(dk._low_res_findings(_gs2, 2)))

print("\n".join("FAIL " + f for f in fails) if fails else "", end="")
print("[test_low_res_image] {}".format("FAILED: {} problem(s)".format(len(fails)) if fails else "ok"))
sys.exit(1 if fails else 0)
