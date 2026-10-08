#!/usr/bin/env python3
"""East-Asian faces follow the SCRIPT: a run the deck's EA face cannot draw (Hangul under a Chinese face) gets that script's face of the same register, in text(), retrofit_ea() and the CJK_NO_EA advice."""
from __future__ import annotations
import sys, tempfile
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
fails: list[str] = []


def check(cond, msg):
    if not cond:
        fails.append(msg)



import deckkit as dk
from pptx.oxml.ns import qn
def ea_of(tb):
    return [(r.text, (r._r.find(".//" + qn("a:ea")).get("typeface") if r._r.find(".//" + qn("a:ea")) is not None else None))
            for p in tb.text_frame.paragraphs for r in p.runs]
HAN, HANGUL_MAC = "Hiragino Sans GB", "Apple SD Gothic Neo"
here = not dk._font_substituted(HAN) and not dk._font_substituted(HANGUL_MAC)
if not here:
    print("  skip per-script EA faces: {} / {} are not installed here".format(HAN, HANGUL_MAC))
else:
    prs = dk.blank_deck(13.333, 7.5)
    old = dk.EAFONT
    dk.EAFONT = HAN
    try:
        s = dk.add_slide(prs)
        tb = dk.text(s, 1, 1, 8, 2, [[("오늘의 순서", 24, dk.DEEP, False, False)], [("城市里的小菜园", 24, dk.DEEP, False, False)],
                                     [("都市の菜園", 24, dk.DEEP, False, False)]])
        got = dict(ea_of(tb))
        check(got["오늘의 순서"] == HANGUL_MAC, "a Korean run under a Chinese EA face that has no Hangul gets a Hangul face: {}".format(got))
        check(got["城市里的小菜园"] == HAN, "a Chinese run keeps the deck's face: {}".format(got))
        check(got["都市の菜園"] == HAN, "kana the Chinese face does draw keeps it (no needless swap): {}".format(got))
    finally:
        dk.EAFONT = old
    # retrofit_ea(prs, <Han face>) on a deck with a bare Korean run stamps a Hangul face there
    prs2 = dk.blank_deck(13.333, 7.5)
    s2 = dk.add_slide(prs2)
    tb2 = s2.shapes.add_textbox(914400, 914400, 914400 * 6, 914400)
    tb2.text_frame.text = "고장 난 물건을 가져오세요"
    dk.retrofit_ea(prs2, HAN, verbose=False)
    check(dict(ea_of(tb2)).get("고장 난 물건을 가져오세요") == HANGUL_MAC, "retrofit_ea gives a bare Korean run a Hangul face: {}".format(ea_of(tb2)))
# the CJK_NO_EA advice names a face for the deck's own script
prs3 = dk.blank_deck(13.333, 7.5)
s3 = dk.add_slide(prs3)
tb3 = s3.shapes.add_textbox(914400, 914400, 914400 * 6, 914400)
tb3.text_frame.text = "고장 난 물건을 가져오세요"
msgs = [f[3] for f in dk.lint_layout(prs3, verbose=False) if f[2] == "CJK_NO_EA"]
check(msgs and ("Apple SD Gothic Neo" in msgs[0] and "Malgun Gothic" in msgs[0]), "a Korean deck is told a Korean face: {}".format(msgs[:1]))

print("\n".join("FAIL " + f for f in fails) if fails else "", end="")
print("[test_ea_script_face] {}".format("FAILED: {} problem(s)".format(len(fails)) if fails else "ok"))
sys.exit(1 if fails else 0)
