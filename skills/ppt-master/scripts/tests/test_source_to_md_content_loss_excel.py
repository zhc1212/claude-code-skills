#!/usr/bin/env python3
"""PPT Master - Spreadsheet content preservation regressions.

Usage: python3 -m unittest tests.test_source_to_md_content_loss_excel
Dependencies: openpyxl
"""

import io
import json
import subprocess
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from datetime import datetime, time
from pathlib import Path

from openpyxl import Workbook

SCRIPTS_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPTS_DIR))
sys.path.insert(0, str(SCRIPTS_DIR / "source_to_md"))
import excel_to_md  # noqa: E402


class ExcelContentLossTests(unittest.TestCase):
    def _convert(self, workbook, **options):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "book.xlsx"
            workbook.save(source)
            err = io.StringIO()
            with redirect_stdout(io.StringIO()), redirect_stderr(err):
                md = excel_to_md.convert_to_markdown(str(source), **options)
            profile = json.loads(source.with_suffix('.conversion_profile.json').read_text(encoding='utf-8'))
            return md, profile, err.getvalue()

    def test_f21_fractional_seconds_are_preserved(self) -> None:
        wb = Workbook()
        wb.active.append(['Datetime', 'Time'])
        wb.active.append([datetime(2026, 10, 5, 12, 34, 56, 789000), time(12, 34, 56, 789000)])
        md, _, _ = self._convert(wb)
        self.assertIn('2026-10-05 12:34:56.789', md)
        self.assertIn('| 12:34:56.789', md)

    def test_f22_percent_and_identifier_formats_are_preserved(self) -> None:
        wb = Workbook()
        sheet = wb.active
        sheet.append(['Percent', 'Identifier'])
        sheet.append([0.15, 123])
        sheet['A2'].number_format = '0%'
        sheet['B2'].number_format = '000000'
        md, _, _ = self._convert(wb)
        self.assertIn('| 15% | 000123 |', md)

    def test_f22_percent_scales_excel_precision_without_binary_noise(self) -> None:
        # openpyxl reads 17-digit floats; Excel itself holds 15 significant digits.
        self.assertEqual(excel_to_md._format_number(0.19497971117842003, "0.00%"), "19.497971117842%")
        self.assertEqual(excel_to_md._format_number(0.1 + 0.2, "0%"), "30%")

    def test_f22_unsupported_format_retains_value_and_format_with_warning(self) -> None:
        wb = Workbook()
        wb.active.append(['Amount'])
        wb.active.append([123.456])
        wb.active['A2'].number_format = '"kg "0.00'
        md, profile, err = self._convert(wb)
        self.assertIn('123.456', md)
        self.assertIn('"kg "0.00', md)
        self.assertIn('(1 cells)', md)
        self.assertTrue(profile['warnings'])
        self.assertIn('number format', err)

    def _hidden_workbook(self):
        wb = Workbook()
        wb.active.append(['Public'])
        for name, state in [('Appendix', 'hidden'), ('Internal', 'veryHidden')]:
            sheet = wb.create_sheet(name)
            sheet.append(['HIDDEN_BODY_' + name])
            sheet.sheet_state = state
        return wb

    def test_s01_hidden_sheets_are_named_and_warned_by_default(self) -> None:
        md, profile, err = self._convert(self._hidden_workbook())
        for name in ['Appendix', 'Internal']:
            self.assertIn(name, md)
            self.assertIn(name, err)
            self.assertIn(name, ' '.join(profile['warnings']))
            self.assertNotIn('HIDDEN_BODY_' + name, md)

    def test_s01_include_hidden_exports_and_marks_sheets(self) -> None:
        md, profile, _ = self._convert(self._hidden_workbook(), include_hidden=True)
        for name in ['Appendix', 'Internal']:
            self.assertIn('## Sheet: ' + name + ' (hidden)', md)
            self.assertIn('HIDDEN_BODY_' + name, md)
        self.assertFalse(profile['warnings'])

    def test_s01_include_hidden_is_forwarded_by_unified_cli(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / 'book.xlsx'
            self._hidden_workbook().save(source)
            proc = subprocess.run([sys.executable, str(SCRIPTS_DIR / 'source_to_md.py'), str(source),
                                   '--include-hidden', '--json'], capture_output=True, text=True)
            self.assertEqual(proc.returncode, 0, proc.stderr)
            self.assertIn('HIDDEN_BODY_Appendix', source.with_suffix('.md').read_text(encoding='utf-8'))

    def test_s02_merged_values_are_copied_with_shared_value_annotation(self) -> None:
        wb = Workbook()
        sheet = wb.active
        sheet.append(['Amount', 'Item'])
        sheet.append([3000000, 'First'])
        sheet['B3'] = 'Second'
        sheet['A2'].number_format = '#,##0'
        sheet.merge_cells('A2:A3')
        md, _, _ = self._convert(wb)
        self.assertIn('| 3000000 | First |', md)
        self.assertIn('| 3000000 | Second |', md)
        self.assertIn('> Merged cells: A2:A3 (shared value)', md)

    def test_display_formats_never_round_or_emit_per_cell_notes(self) -> None:
        wb = Workbook()
        sheet = wb.active
        sheet.append(['Amount', 'Format'])
        formats = ['0', '0.00', '#,##0', '#,##0.00', '0.00_', '0_', '"USD "0.00',
                   '_(* #,##0.00_);_(* (#,##0.00);_(* "-"??_);_(@_)', '@']
        for fmt in formats:
            sheet.append([74083.6551, fmt])
            sheet.cell(sheet.max_row, 1).number_format = fmt
        md, profile, err = self._convert(wb)
        self.assertEqual(md.count('| 74083.6551 |'), len(formats))
        self.assertFalse(profile['warnings'])
        self.assertEqual(err, '')
        self.assertNotIn('number format', md)

    def test_percentages_keep_full_precision(self) -> None:
        wb = Workbook()
        sheet = wb.active
        sheet.append(['Percent'])
        for value in (0.036, 0.551282051282051, -0.036, 1.00001):
            sheet.append([value])
            sheet.cell(sheet.max_row, 1).number_format = '0.00%'
        md, _, _ = self._convert(wb)
        for value in ('3.6%', '55.1282051282051%', '-3.6%', '100.001%'):
            self.assertIn('| ' + value + ' |', md)

    def test_identifier_does_not_round_fractional_values(self) -> None:
        wb = Workbook()
        wb.active.append(['Identifier'])
        for value in (123.456, -123):
            wb.active.append([value])
            wb.active.cell(wb.active.max_row, 1).number_format = '000000'
        md, _, _ = self._convert(wb)
        self.assertIn('| 123.456 |', md)
        self.assertIn('| -000123 |', md)

    def test_unknown_semantic_formats_are_summarized_once_per_sheet(self) -> None:
        wb = Workbook()
        wb.active.append(['Amount'])
        for fmt in ('"kg "0.00', '"kg "0.00', '0.00"m"'):
            wb.active.append([123.456])
            wb.active.cell(wb.active.max_row, 1).number_format = fmt
        md, profile, err = self._convert(wb)
        self.assertEqual(md.count('unsupported number formats:'), 1)
        self.assertIn('(2 cells)', md)
        self.assertIn('(1 cells)', md)
        self.assertEqual(len(profile['warnings']), 1)
        self.assertEqual(err.count('[WARN]'), 1)

    def test_text_merges_do_not_emit_annotations(self) -> None:
        wb = Workbook()
        wb.active.append(['Title', None])
        wb.active.merge_cells('A1:B1')
        wb.active.append(['Amount', 'Item'])
        wb.active.append([12, 'Data'])
        md, _, _ = self._convert(wb)
        self.assertIn('Title', md)
        self.assertNotIn('Merged cells:', md)

    def test_percent_sections_and_color_keep_precision(self) -> None:
        wb = Workbook()
        sheet = wb.active
        sheet.append(['Percent'])
        for value, fmt in ((0.036, r'\+0%;\-0%'), (-0.036, r'\+0%;\-0%'),
                           (0.551282051282051, '0.00%;[Red](0.00%)')):
            sheet.append([value])
            sheet.cell(sheet.max_row, 1).number_format = fmt
        md, profile, _ = self._convert(wb)
        for value in ('3.6%', '-3.6%', '55.1282051282051%'):
            self.assertIn('| ' + value + ' |', md)
        self.assertFalse(profile['warnings'])
