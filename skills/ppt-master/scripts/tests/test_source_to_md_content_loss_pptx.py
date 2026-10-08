#!/usr/bin/env python3
"""PPT Master - PowerPoint content preservation regressions.

Usage: python3 -m unittest tests.test_source_to_md_content_loss_pptx
Dependencies: python-pptx
"""

import io
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path

from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches

SCRIPTS_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPTS_DIR))
sys.path.insert(0, str(SCRIPTS_DIR / 'source_to_md'))
import ppt_to_md  # noqa: E402


class PptxContentLossTests(unittest.TestCase):
    def _presentation(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
        return prs, shape

    def _convert(self, prs):
        with tempfile.TemporaryDirectory() as tmp, redirect_stdout(io.StringIO()):
            path = Path(tmp) / 'deck.pptx'
            prs.save(path)
            return ppt_to_md.convert_presentation_to_markdown(str(path), str(path.with_suffix('.md')))

    def test_f05_field_text_keeps_xml_order(self) -> None:
        prs, shape = self._presentation()
        p = shape.text_frame.paragraphs[0]
        p.text = 'Date: '
        field = OxmlElement('a:fld')
        field.set('id', '{11111111-1111-1111-1111-111111111111}')
        field.set('type', 'datetime1')
        text = OxmlElement('a:t')
        text.text = '2026-10-05'
        field.append(text)
        p._p.append(field)
        p.add_line_break()
        p.add_run().text = 'After'
        md = self._convert(prs)
        self.assertIn('Date: 2026-10-05\nAfter', md)

    def test_f05_field_hyperlink_is_preserved(self) -> None:
        from pptx.text.text import _Run

        prs, shape = self._presentation()
        p = shape.text_frame.paragraphs[0]
        field = OxmlElement('a:fld')
        field.set('id', '{11111111-1111-1111-1111-111111111111}')
        field.set('type', 'datetime1')
        text = OxmlElement('a:t')
        text.text = '2026-10-05'
        field.append(text)
        p._p.append(field)
        _Run(field, p).hyperlink.address = 'https://example.invalid/date'
        self.assertIn('[2026-10-05](https://example.invalid/date)', self._convert(prs))

    def test_plain_table_cell_does_not_require_shape_styles(self) -> None:
        prs, shape = self._presentation()
        slide = prs.slides[0]
        table = slide.shapes.add_table(1, 1, Inches(1), Inches(3), Inches(4), Inches(1)).table
        table.cell(0, 0).text = 'Cell body'
        self.assertEqual(ppt_to_md.text_frame_to_markdown(table.cell(0, 0).text_frame), 'Cell body')
        self.assertIn('Cell body', self._convert(prs))

    def test_f25_explicit_start_and_continuation_are_preserved(self) -> None:
        prs, shape = self._presentation()
        for i, label in enumerate(['Fifth', 'Sixth', 'Seventh']):
            p = shape.text_frame.paragraphs[0] if i == 0 else shape.text_frame.add_paragraph()
            p.text = label
            bullet = OxmlElement('a:buAutoNum')
            bullet.set('type', 'arabicPeriod')
            if i < 2:
                bullet.set('startAt', str(5 + i))
            p._p.get_or_add_pPr().append(bullet)
        md = self._convert(prs)
        for text in ('5. Fifth', '6. Sixth', '7. Seventh'):
            self.assertIn(text, md)

    def test_f25_inherited_list_style_and_explicit_none(self) -> None:
        prs, shape = self._presentation()
        lst = shape.text_frame._txBody.find(ppt_to_md.qn('a:lstStyle'))
        level = OxmlElement('a:lvl1pPr')
        bullet = OxmlElement('a:buAutoNum')
        bullet.set('type', 'alphaLcParenR')
        bullet.set('startAt', '5')
        level.append(bullet)
        lst.append(level)
        for i, label in enumerate(['Fifth', 'Sixth', 'Plain']):
            p = shape.text_frame.paragraphs[0] if i == 0 else shape.text_frame.add_paragraph()
            p.text = label
            if i == 2:
                p._p.get_or_add_pPr().append(OxmlElement('a:buNone'))
        md = self._convert(prs)
        self.assertIn('e) Fifth', md)
        self.assertIn('f) Sixth', md)
        self.assertIn('\n\nPlain', md)

    def test_f25_numbering_jumps_keep_literal_labels(self) -> None:
        prs, shape = self._presentation()
        for i, number in enumerate((5, 9)):
            p = shape.text_frame.paragraphs[0] if i == 0 else shape.text_frame.add_paragraph()
            p.text = f'Item {number}'
            bullet = OxmlElement('a:buAutoNum')
            bullet.set('type', 'arabicPeriod')
            bullet.set('startAt', str(number))
            p._p.get_or_add_pPr().append(bullet)
        md = self._convert(prs)
        self.assertIn(r'5\. Item 5', md)
        self.assertIn(r'9\. Item 9', md)

    def test_f25_layout_inherited_numbering(self) -> None:
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        body = slide.placeholders[1]
        layout_body = slide.slide_layout.placeholders.get(body.placeholder_format.idx)
        p = layout_body.text_frame.paragraphs[0]
        props = p._p.get_or_add_pPr()
        for child in list(props):
            props.remove(child)
        bullet = OxmlElement('a:buAutoNum')
        bullet.set('type', 'arabicPeriod')
        bullet.set('startAt', '5')
        props.append(bullet)
        body.text = 'Fifth\nSixth'
        md = self._convert(prs)
        self.assertIn('5. Fifth', md)
        self.assertIn('6. Sixth', md)

    def test_f27_plain_paragraphs_do_not_acquire_bullets(self) -> None:
        prs, shape = self._presentation()
        shape.text = 'First paragraph\nSecond paragraph'
        md = self._convert(prs)
        self.assertIn('First paragraph\n\nSecond paragraph', md)
        self.assertNotIn('- First paragraph', md)

    def test_f27_real_bullets_are_retained_per_paragraph(self) -> None:
        prs, shape = self._presentation()
        p = shape.text_frame.paragraphs[0]
        p.text = 'List item'
        bullet = OxmlElement('a:buChar')
        bullet.set('char', '•')
        p._p.get_or_add_pPr().append(bullet)
        shape.text_frame.add_paragraph().text = 'Plain paragraph'
        md = self._convert(prs)
        self.assertIn('- List item\n\nPlain paragraph', md)

    def test_literal_numbered_paragraphs_have_no_blank_separator(self) -> None:
        prs, shape = self._presentation()
        shape.text = '1. First\n2. Second\n3) Third'
        md = self._convert(prs)
        self.assertIn('1. First\n2. Second\n3) Third', md)

    def test_table_paragraphs_use_one_break(self) -> None:
        prs, _ = self._presentation()
        table = prs.slides[0].shapes.add_table(1, 1, Inches(1), Inches(3), Inches(4), Inches(1)).table
        table.cell(0, 0).text = 'First\nSecond'
        md = self._convert(prs)
        self.assertIn('| First<br>Second |', md)
        self.assertNotIn('<br><br>', md)
