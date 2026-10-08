#!/usr/bin/env python3
"""PPT Master - DOCX content preservation regressions.

Usage: python3 -m unittest tests.test_source_to_md_content_loss_docx
Dependencies: python-docx, mammoth
"""

import io
import json
import sys
import tempfile
import unittest
import zipfile
from contextlib import redirect_stdout
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.opc.constants import RELATIONSHIP_TYPE as RT

SCRIPTS_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPTS_DIR))
sys.path.insert(0, str(SCRIPTS_DIR / "source_to_md"))
import doc_to_md  # noqa: E402


class DocxContentLossTests(unittest.TestCase):
    def _convert(self, document, patch_archive=None):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.docx"
            if patch_archive:
                buf = io.BytesIO()
                document.save(buf)
                patch_archive(buf, source)
            else:
                document.save(source)
            with redirect_stdout(io.StringIO()):
                markdown = doc_to_md.convert_to_markdown(str(source))
            profile = json.loads(source.with_suffix(".conversion_profile.json").read_text(encoding="utf-8"))
            return markdown, profile

    def test_f01_nested_table_body_is_preserved(self) -> None:
        doc = Document()
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "Outer body"
        cell.add_table(rows=1, cols=1).cell(0, 0).text = "NESTED_SENTINEL_3000000"
        md, _ = self._convert(doc)
        self.assertIn("Outer body", md)
        self.assertIn("NESTED_SENTINEL_3000000", md)

    def test_f02_content_control_display_text_is_preserved(self) -> None:
        doc = Document()
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "Before"
        cell._tc.append(parse_xml('<w:sdt ' + nsdecls('w') + '><w:sdtContent><w:p><w:r>'
                                 '<w:t>CONTROL_BODY</w:t></w:r></w:p></w:sdtContent></w:sdt>'))
        md, profile = self._convert(doc)
        self.assertIn("CONTROL_BODY", md)
        self.assertTrue(any("content control" in w.lower() for w in profile["warnings"]))

    def test_f03_table_endnote_reference_and_body_are_preserved(self) -> None:
        doc = Document()
        p = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
        p.add_run("Reference")
        p._p.append(parse_xml('<w:r ' + nsdecls('w') + '><w:endnoteReference w:id="1"/></w:r>'))

        def patch_archive(buf, source):
            with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(source, "w") as zout:
                for item in zin.infolist():
                    data = zin.read(item.filename)
                    if item.filename == "word/_rels/document.xml.rels":
                        root = ET.fromstring(data)
                        ET.SubElement(root, '{http://schemas.openxmlformats.org/package/2006/relationships}'
                                      'Relationship', Id="rIdEndnotes", Type=RT.ENDNOTES, Target="endnotes.xml")
                        data = ET.tostring(root)
                    if item.filename == "[Content_Types].xml":
                        root = ET.fromstring(data)
                        ET.SubElement(root, '{http://schemas.openxmlformats.org/package/2006/content-types}'
                                      'Override', PartName="/word/endnotes.xml", ContentType=
                                      "application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml")
                        data = ET.tostring(root)
                    zout.writestr(item, data)
                zout.writestr("word/endnotes.xml", '<w:endnotes ' + nsdecls('w') + '><w:endnote w:id="1">'
                              '<w:p><w:r><w:t>ENDNOTE_BODY</w:t></w:r></w:p></w:endnote></w:endnotes>')

        md, _ = self._convert(doc, patch_archive)
        self.assertIn("[^endnote-1]", md)
        self.assertIn("[^endnote-1]: ENDNOTE_BODY", md)

    def test_f09_table_hyperlink_target_is_preserved(self) -> None:
        doc = Document()
        p = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
        rid = p.part.relate_to("https://example.invalid/evidence", RT.HYPERLINK, is_external=True)
        p._p.append(parse_xml('<w:hyperlink ' + nsdecls('w', 'r') + ' r:id="' + rid + '">'
                             '<w:r><w:t>Evidence</w:t></w:r></w:hyperlink>'))
        md, _ = self._convert(doc)
        self.assertIn("| [Evidence](https://example.invalid/evidence) |", md)
        self.assertIn("| --- |", md)

    def _math_document(self, xml):
        doc = Document()
        doc.add_paragraph()._p.append(parse_xml('<m:oMath ' + nsdecls('m') + '>' + xml + '</m:oMath>'))
        return doc

    def test_f23_delimiter_operands_keep_separator(self) -> None:
        for properties, separator in [('<m:dPr><m:sepChr m:val="|"/></m:dPr>', '|'), ('', '|'),
                                      ('<m:dPr><m:sepChr m:val=";"/></m:dPr>', ';')]:
            with self.subTest(separator=separator, properties=properties):
                md, _ = self._convert(self._math_document('<m:d>' + properties +
                    '<m:e><m:r><m:t>x</m:t></m:r></m:e><m:e><m:r><m:t>y</m:t></m:r></m:e></m:d>'))
                self.assertIn('x' + separator + 'y', md)

    def test_f24_group_character_and_position_are_preserved(self) -> None:
        for char, pos, expected in [('⏞', 'top', r'\overbrace{x+y}'),
                                    ('⏟', 'bot', r'\underbrace{x+y}')]:
            with self.subTest(char=char):
                md, _ = self._convert(self._math_document('<m:groupChr><m:groupChrPr><m:chr m:val="' + char +
                    '"/><m:pos m:val="' + pos + '"/></m:groupChrPr><m:e><m:r><m:t>x+y</m:t></m:r>'
                    '</m:e></m:groupChr>'))
                self.assertIn(expected, md)

    def test_f24_unknown_group_character_is_retained_with_warning(self) -> None:
        md, profile = self._convert(self._math_document('<m:groupChr><m:groupChrPr><m:chr m:val="?"/>'
            '</m:groupChrPr><m:e><m:r><m:t>x</m:t></m:r></m:e></m:groupChr>'))
        self.assertIn('?', md)
        self.assertTrue(any('group character' in w for w in profile['warnings']))
