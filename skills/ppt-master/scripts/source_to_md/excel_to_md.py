#!/usr/bin/env python3
"""
Excel to Markdown Converter

Supported formats:
    .xlsx   Excel workbook
    .xlsm   Excel macro-enabled workbook

Unsupported by default:
    .xls    Legacy binary Excel format; resave as .xlsx first

All paths produce the same output convention:
    <input>.md                     Markdown file
"""

import argparse
import re
import sys
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any

_SCRIPTS_DIR = Path(__file__).resolve().parents[1]
if str(_SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS_DIR))

from console_encoding import configure_utf8_stdio  # noqa: E402
from _batch import run_path_batch  # noqa: E402
from _conversion_profile import write_conversion_profile_best_effort  # noqa: E402

configure_utf8_stdio()


# ─────────────────────────────────────────────────────────────
# Format registry
# ─────────────────────────────────────────────────────────────

EXCEL_FORMATS = {".xlsx", ".xlsm"}
LEGACY_EXCEL_FORMATS = {".xls"}


# ─────────────────────────────────────────────────────────────
# Shared helpers
# ─────────────────────────────────────────────────────────────

def _format_size(size: int) -> str:
    for unit in ("B", "KB", "MB"):
        if size < 1024:
            return f"{size:.0f} {unit}"
        size /= 1024
    return f"{size:.1f} GB"


def _report_result(out_file: Path) -> None:
    size = out_file.stat().st_size
    print(f"[OK] Saved Markdown to: {out_file} ({_format_size(size)})")


def _is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _markdown_escape(value: str) -> str:
    value = value.replace("\\", "\\\\")
    value = value.replace("|", "\\|")
    value = value.replace("\r\n", "\n").replace("\r", "\n")
    return re.sub(r"\s*\n\s*", "<br>", value).strip()


def _format_number(value: int | float, number_format: str) -> str | None:
    """Add percentage/identifier meaning without rounding the stored value."""
    sections = re.split(r';(?=(?:[^"]*"[^"]*")*[^"]*$)', number_format)
    section = sections[1] if value < 0 and len(sections) > 1 else sections[0]
    section = _strip_display_tokens(section)
    if re.fullmatch(r"[+\-()]?(?:0+|#,##0)(?:\.[0#]+)?%\)?", section):
        # Match the plain-number path: Excel holds 15 significant digits, so
        # scale that value rather than the 17-digit binary repr.
        number = Decimal(f"{value:.15g}") * 100
        if number.is_finite():
            rendered = format(number, "f")
            if "." in rendered:
                rendered = rendered.rstrip("0").rstrip(".")
            return rendered + "%"
    if re.fullmatch(r"0+", section) and value == int(value):
        return ("-" if value < 0 else "") + str(abs(int(value))).zfill(len(section))
    return None


def _strip_display_tokens(number_format: str) -> str:
    """Remove accounting decoration while retaining units and scale tokens."""
    # Accounting spacing, colors, currency locales, and currency labels do not
    # change a raw number's meaning. Other literal units or scaling can do so.
    cleaned = re.sub(r"_.?|\*.", "", number_format)
    cleaned = re.sub(r"\[(?:Black|Blue|Cyan|Green|Magenta|Red|White|Yellow|Color\d+|\$[^\]]*)\]",
                     "", cleaned, flags=re.I)
    cleaned = re.sub(r'"(?:USD|EUR|GBP|CNY|RMB|JPY|HKD|[$€£¥￥\s-])*"', "", cleaned)
    cleaned = re.sub(r"\\([$€£¥￥ ()+-])", r"\1", cleaned)
    return cleaned.strip()


def _format_needs_warning(number_format: str) -> bool:
    """Identify unhandled semantic formats, excluding ordinary display styling."""
    if number_format in {"General", "@"}:
        return False
    cleaned = _strip_display_tokens(number_format)
    return (not re.fullmatch(r"[0-9#?.,Ee+\-;/()@$€£¥￥\s]*", cleaned)
            or bool(re.search(r",\s*(?:;|$)", cleaned)))


def _format_cell_value(value: Any, number_format: str = "General") -> str:
    if _is_empty(value):
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if _is_numeric_value(value):
        formatted = _format_number(value, number_format)
        if formatted is not None:
            return formatted
    if isinstance(value, float):
        # Excel shows 15 significant digits; that keeps 1234567.89 exact while
        # 0.1 + 0.2 reads as 0.3 rather than its binary expansion.
        return _markdown_escape(repr(float(f"{value:.15g}")).removesuffix(".0"))
    return _markdown_escape(str(value))


def _is_numeric_value(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _sheet_state_label(sheet_state: str) -> str:
    if sheet_state == "visible":
        return "visible"
    return sheet_state or "unknown"


# ─────────────────────────────────────────────────────────────
# Worksheet extraction
# ─────────────────────────────────────────────────────────────

def _merged_value_map(worksheet) -> dict[tuple[int, int], Any]:
    """Return propagated values for merged cells, keyed by (row, column).

    Merged regions whose top-left cell is empty are intentionally skipped.
    These are typically formatting-only ranges that carry no textual content.
    """
    merged_values: dict[tuple[int, int], Any] = {}
    for merged_range in worksheet.merged_cells.ranges:
        value = worksheet.cell(merged_range.min_row, merged_range.min_col).value
        if _is_empty(value):
            continue
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for col in range(merged_range.min_col, merged_range.max_col + 1):
                merged_values[(row, col)] = value
    return merged_values


def _cell_value(worksheet, row: int, col: int, merged_values: dict[tuple[int, int], Any]) -> Any:
    value = worksheet.cell(row, col).value
    if _is_empty(value):
        return merged_values.get((row, col), value)
    return value


def _content_bounds(worksheet, merged_values: dict[tuple[int, int], Any]) -> tuple[int, int, int, int] | None:
    min_row = min_col = None
    max_row = max_col = None

    for row in worksheet.iter_rows():
        for cell in row:
            if _is_empty(cell.value):
                continue
            min_row = cell.row if min_row is None else min(min_row, cell.row)
            max_row = cell.row if max_row is None else max(max_row, cell.row)
            min_col = cell.column if min_col is None else min(min_col, cell.column)
            max_col = cell.column if max_col is None else max(max_col, cell.column)

    for row, col in merged_values:
        min_row = row if min_row is None else min(min_row, row)
        max_row = row if max_row is None else max(max_row, row)
        min_col = col if min_col is None else min(min_col, col)
        max_col = col if max_col is None else max(max_col, col)

    if min_row is None or min_col is None or max_row is None or max_col is None:
        return None
    return min_row, min_col, max_row, max_col


def _trim_trailing_empty_cells(row: list[Any]) -> list[Any]:
    trimmed = list(row)
    while trimmed and _is_empty(trimmed[-1]):
        trimmed.pop()
    return trimmed


def _extract_rows(
    worksheet,
    bounds: tuple[int, int, int, int],
    merged_values: dict[tuple[int, int], Any],
    max_rows: int,
    max_cols: int,
) -> tuple[list[list[Any]], bool, bool]:
    min_row, min_col, max_row, max_col = bounds

    row_limit = max_row
    col_limit = max_col
    rows_truncated = False
    cols_truncated = False

    if max_rows > 0 and (max_row - min_row + 1) > max_rows:
        row_limit = min_row + max_rows - 1
        rows_truncated = True
    if max_cols > 0 and (max_col - min_col + 1) > max_cols:
        col_limit = min_col + max_cols - 1
        cols_truncated = True

    rows: list[list[Any]] = []
    width = 0
    for row_index in range(min_row, row_limit + 1):
        row = [
            _cell_value(worksheet, row_index, col_index, merged_values)
            for col_index in range(min_col, col_limit + 1)
        ]
        row = _trim_trailing_empty_cells(row)
        width = max(width, len(row))
        rows.append(row)

    if width == 0:
        return [], rows_truncated, cols_truncated

    normalized_rows = [row + [""] * (width - len(row)) for row in rows]
    return normalized_rows, rows_truncated, cols_truncated


def _is_caption_row(
    row: list[Any],
    row_index: int,
    min_col: int,
    merged_values: dict[tuple[int, int], Any],
) -> bool:
    """Return whether a leading row is a sheet title or source note, not a header.

    Public datasets (World Bank, OECD, Eurostat exports) put a free-text
    source line above the real header; it fills only the first cell or one
    merged band, so every other column is empty or repeats the same value.
    """
    values = [value for value in row if not _is_empty(value)]
    if len(values) != 1 and len({str(value) for value in values}) != 1:
        return False
    if len(values) == 1:
        return len(row) > 1
    return (row_index, min_col) in merged_values


def _split_caption_rows(
    rows: list[list[Any]],
    min_row: int,
    min_col: int,
    merged_values: dict[tuple[int, int], Any],
) -> tuple[list[str], list[list[Any]]]:
    """Peel title/source rows off the top so the first table row is the header."""
    captions: list[str] = []
    while len(rows) > 1 and _is_caption_row(rows[0], min_row + len(captions), min_col, merged_values):
        remaining = rows[1:]
        if not any(len([v for v in row if not _is_empty(v)]) >= 2 for row in remaining):
            break
        captions.append(_format_cell_value(next(value for value in rows[0] if not _is_empty(value))))
        rows = remaining
    return captions, rows


def _column_alignments(rows: list[list[Any]]) -> list[str]:
    if not rows:
        return []

    width = len(rows[0])
    alignments: list[str] = []
    data_rows = rows[1:] if len(rows) > 1 else rows
    for col_index in range(width):
        values = [row[col_index] for row in data_rows if not _is_empty(row[col_index])]
        if values and all(_is_numeric_value(value) for value in values):
            alignments.append("---:")
        else:
            alignments.append("---")
    return alignments


def _rows_to_markdown_table(
    rows: list[list[Any]], number_formats: list[list[str]] | None = None,
) -> str:
    if not rows:
        return "_No tabular content found._"

    formatted_rows = [
        [_format_cell_value(value, number_formats[i][j] if number_formats else "General")
         for j, value in enumerate(row)]
        for i, row in enumerate(rows)
    ]
    width = len(formatted_rows[0])
    separator = _column_alignments(rows)
    lines = [
        "| " + " | ".join(formatted_rows[0]) + " |",
        "| " + " | ".join(separator or ["---"] * width) + " |",
    ]

    for row in formatted_rows[1:]:
        lines.append("| " + " | ".join(row) + " |")

    return "\n".join(lines)


# ─────────────────────────────────────────────────────────────
# Excel → Markdown
# ─────────────────────────────────────────────────────────────

def _convert_excel(
    input_file: Path, out_file: Path, max_rows: int, max_cols: int,
    include_hidden: bool = False, warnings: list[str] | None = None,
) -> str:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("[ERROR] openpyxl not installed. Run: pip install openpyxl")
        return ""

    workbook = load_workbook(input_file, data_only=True, read_only=False)
    visible_sheets = [sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
    hidden_sheets = [sheet for sheet in workbook.worksheets if sheet.sheet_state != "visible"]
    selected_sheets = workbook.worksheets if include_hidden else visible_sheets

    lines: list[str] = [
        f"# Spreadsheet Source: {input_file.name}",
        "",
        "## Workbook Summary",
        "",
        f"- Sheets: {len(workbook.worksheets)}",
        f"- Visible sheets: {', '.join(sheet.title for sheet in visible_sheets) or 'None'}",
        "",
        "> Note: Formula cells are exported as cached values. This converter does not recalculate formulas.",
        "",
    ]

    if hidden_sheets and not include_hidden:
        warning = "Skipped hidden sheets: " + ", ".join(sheet.title for sheet in hidden_sheets)
        warning += ". Use --include-hidden to export them."
        lines.extend([f"> Warning: {warning}", ""])
        print(f"[WARN] {warning}", file=sys.stderr)
        if warnings is not None:
            warnings.append(warning)

    if not selected_sheets:
        lines.extend(["_No visible sheets found._", ""])

    for worksheet in selected_sheets:
        merged_values = _merged_value_map(worksheet)
        bounds = _content_bounds(worksheet, merged_values)

        lines.extend([
            f"## Sheet: {worksheet.title}" + (" (hidden)" if worksheet.sheet_state != "visible" else ""),
            "",
            f"- State: {_sheet_state_label(worksheet.sheet_state)}",
        ])

        if bounds is None:
            lines.extend(["", "_No content found._", ""])
            continue

        min_row, min_col, max_row, max_col = bounds
        used_range = (
            f"{get_column_letter(min_col)}{min_row}:"
            f"{get_column_letter(max_col)}{max_row}"
        )
        rows, rows_truncated, cols_truncated = _extract_rows(
            worksheet,
            bounds,
            merged_values,
            max_rows=max_rows,
            max_cols=max_cols,
        )

        lines.extend([
            f"- Used range: {used_range}",
            f"- Rows: {max_row - min_row + 1}",
            f"- Columns: {max_col - min_col + 1}",
            "",
        ])

        # Propagated merged values must use the anchor's display format too.
        format_anchors = {
            (row, col): (region.min_row, region.min_col)
            for region in worksheet.merged_cells.ranges
            for row in range(region.min_row, region.max_row + 1)
            for col in range(region.min_col, region.max_col + 1)
            if (row, col) in merged_values
        }
        number_formats = []
        unsupported_formats: dict[str, int] = {}
        for i, row in enumerate(rows, min_row):
            formats = []
            for j, value in enumerate(row, min_col):
                anchor = format_anchors.get((i, j), (i, j))
                cell = worksheet.cell(*anchor)
                fmt = cell.number_format
                formats.append(fmt)
                if (_is_numeric_value(value) and anchor == (i, j)
                        and _format_number(value, fmt) is None and _format_needs_warning(fmt)):
                    unsupported_formats[fmt] = unsupported_formats.get(fmt, 0) + 1
            number_formats.append(formats)

        original_rows = rows
        captions, rows = _split_caption_rows(rows, min_row, min_col, merged_values)
        for i in range(len(captions)):
            j = next(j for j, value in enumerate(original_rows[i]) if not _is_empty(value))
            captions[i] = _format_cell_value(original_rows[i][j], number_formats[i][j])
        number_formats = number_formats[len(captions):]
        for caption in captions:
            lines.extend([f"> {caption}", ""])

        if rows_truncated or cols_truncated:
            limit_notes = []
            if rows_truncated:
                limit_notes.append(f"rows limited to {max_rows}")
            if cols_truncated:
                limit_notes.append(f"columns limited to {max_cols}")
            lines.extend([f"> Truncated: {', '.join(limit_notes)}.", ""])

        lines.extend([_rows_to_markdown_table(rows, number_formats), ""])
        for region in sorted(worksheet.merged_cells.ranges, key=lambda r: (r.min_row, r.min_col)):
            value = worksheet.cell(region.min_row, region.min_col).value
            if _is_numeric_value(value):
                lines.extend([f"> Merged cells: {region} (shared value)", ""])
        if unsupported_formats:
            formats = "; ".join(f"{fmt!r} ({count} cells)" for fmt, count in unsupported_formats.items())
            warning = f"{worksheet.title}: unsupported number formats: {formats}; raw values retained."
            print(f"[WARN] {warning}", file=sys.stderr)
            if warnings is not None:
                warnings.append(warning)
            lines.extend([f"> {_markdown_escape(warning)}", ""])

    workbook.close()
    markdown = "\n".join(lines).rstrip() + "\n"
    out_file.write_text(markdown, encoding="utf-8")
    _report_result(out_file)
    return markdown


# ─────────────────────────────────────────────────────────────
# Dispatcher
# ─────────────────────────────────────────────────────────────

def convert_to_markdown(
    input_path: str,
    output_path: str | None = None,
    max_rows: int = 0,
    max_cols: int = 0,
    include_hidden: bool = False,
) -> str:
    input_file = Path(input_path)
    if not input_file.exists():
        print(f"[ERROR] File not found: {input_path}")
        return ""

    suffix = input_file.suffix.lower()
    if suffix in LEGACY_EXCEL_FORMATS:
        print("[ERROR] Unsupported legacy Excel format: .xls")
        print("   Please resave the workbook as .xlsx and run this converter again.")
        return ""
    if suffix not in EXCEL_FORMATS:
        supported = ", ".join(sorted(EXCEL_FORMATS))
        print(f"[ERROR] Unsupported format: {suffix}")
        print(f"   Supported: {supported}")
        return ""

    if max_rows < 0 or max_cols < 0:
        print("[ERROR] --max-rows and --max-cols must be zero or positive integers")
        return ""

    out_file = Path(output_path) if output_path else input_file.with_suffix(".md")
    out_file.parent.mkdir(parents=True, exist_ok=True)

    print(f"[INFO] Converting Excel workbook: {input_file.name}")
    warnings: list[str] = []
    markdown = _convert_excel(
        input_file, out_file, max_rows=max_rows, max_cols=max_cols,
        include_hidden=include_hidden, warnings=warnings,
    )
    if markdown:
        profile_path = write_conversion_profile_best_effort(
            input_path=str(input_file),
            markdown_path=out_file,
            converter="excel_to_md.py",
            conversion_type=suffix.lstrip("."),
            warnings=warnings,
        )
        if profile_path:
            print(f"   Wrote conversion profile -> {profile_path}")
    return markdown


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Convert Excel workbooks to Markdown",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  python excel_to_md.py report.xlsx
  python excel_to_md.py report.xlsx budget.xlsm
  python excel_to_md.py ./workbooks -o ./markdown
  python excel_to_md.py report.xlsx -o output.md
  python excel_to_md.py report.xlsm --max-rows 200 --max-cols 40

Supported formats:
  .xlsx  .xlsm

Unsupported by default:
  .xls   Resave as .xlsx first
        """,
    )
    parser.add_argument("inputs", nargs="+", help="Input Excel workbook(s) or directories")
    parser.add_argument(
        "-o",
        "--output",
        help="Output Markdown file for one input, or output directory for multiple inputs/directories",
    )
    parser.add_argument(
        "--max-rows",
        type=int,
        default=0,
        help="Maximum rows per sheet to export (0 = no limit)",
    )
    parser.add_argument(
        "--max-cols",
        type=int,
        default=0,
        help="Maximum columns per sheet to export (0 = no limit)",
    )
    parser.add_argument("--include-hidden", action="store_true", help="Export hidden worksheets as well")
    args = parser.parse_args()

    return run_path_batch(
        args.inputs,
        EXCEL_FORMATS | LEGACY_EXCEL_FORMATS,
        args.output,
        lambda source, output: bool(
            convert_to_markdown(
                str(source),
                str(output),
                max_rows=args.max_rows,
                max_cols=args.max_cols,
                include_hidden=args.include_hidden,
            )
        ),
    )


if __name__ == "__main__":
    raise SystemExit(main())
